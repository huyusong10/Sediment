from __future__ import annotations

import ast
import json
import re
import sys
import textwrap
from collections import Counter
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Any, Iterable

import yaml

SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR.parent.parent
SRC_ROOT = PROJECT_ROOT / "src"
if str(SRC_ROOT) not in sys.path:
    sys.path.insert(0, str(SRC_ROOT))

from sediment.kb import inventory as kb_inventory

try:
    from docx import Document as DocxDocument
except ImportError:  # pragma: no cover - dependency is optional at import time
    DocxDocument = None


_GENERIC_TERMS = {
    "说明",
    "描述",
    "定义",
    "版本",
    "状态",
    "流程",
    "规则",
    "配置",
    "当前",
    "目录",
    "注意",
    "步骤",
    "文档",
    "概述",
    "总则",
    "检查",
    "检测",
    "评估",
    "记录",
    "权限",
    "故障",
    "应急",
    "维护团队",
    "审批人",
    "生效日期",
    "文档编号",
    "适用范围",
    "适用对象",
    "负责人",
    "发布时间",
    "更新时间",
    "以下",
    "以内",
    "建议",
    "执行",
    "发现",
    "判断",
    "区域",
    "小时",
    "合格",
    "危险",
    "响应",
    "事件编号",
    "参会角色",
    "发布单位",
    "发布日期",
    "主要原因",
    "容量",
    "平台建设背景",
    "平台里程碑",
    "年度审计结果",
    "年度核心",
    "升级计划与预算",
    "实操记录",
    "夜班值班交接材料",
    "了解行业中的主要角色和团队",
    "技术",
    "报告",
    "调查",
    "复盘",
    "异常",
    "消息类型",
    "指标",
    "核心指标",
    "数据质量",
    "年度核心指标",
    "总览",
    "统计",
    "基本信息",
    "生命周期",
}
_SUMMARY_PATTERNS = (
    "{term}是",
    "{term}既是",
    "{term}指",
    "{term}用于",
    "{term}负责",
    "{term}表示",
    "{term}衡量",
    "{term}属于",
    "{term}作为",
    "{term}提供",
)
_RELATION_SPLIT = re.compile(r"[，。；：:\n]")
_TERM_PATTERN = re.compile(r"[A-Za-z][A-Za-z0-9_.-]*|[\u4e00-\u9fff]{2,16}")
_SENTENCE_PATTERN = re.compile(r"(?<=[。！？!?；;\n])")
_DEFINITION_MARKERS = ("是", "指", "用于", "负责", "表示", "衡量", "属于", "作为", "算法", "流程", "装置", "设备", "团队", "角色", "系统", "协议", "策略", "模式", "状态", "工具", "技术", "操作")
_DIRECT_DEFINITION_MARKERS = ("是", "指", "用于", "负责", "表示", "衡量", "属于", "作为")
_CANONICAL_BARE_TERM_PATTERNS = (
    (re.compile(r"([\u4e00-\u9fff]{2,8})阶段"), "stage"),
    (re.compile(r"([\u4e00-\u9fff]{2,8})追踪(?:记录|标识)"), "trace"),
    (re.compile(r"([\u4e00-\u9fff]{2,8})频率设定与校准"), "calibration"),
    (re.compile(r"掌握([\u4e00-\u9fff]{2,8})标准操作流程"), "sop"),
    (re.compile(r"([\u4e00-\u9fff]{2,8})标准操作流程"), "sop"),
    (re.compile(r"([\u4e00-\u9fff]{2,8})操作流程"), "sop"),
    (re.compile(r"([\u4e00-\u9fff]{2,8})模式"), "mode"),
    (re.compile(r"([\u4e00-\u9fff]{2,8})数据质量"), "metric"),
    (re.compile(r"与([\u4e00-\u9fff]{2,8})行动"), "initiative"),
    (re.compile(r"([\u4e00-\u9fff]{2,8})专项"), "initiative"),
)
_QUERYABLE_ARTIFACT_ALIAS_MARKERS = (
    ".json",
    ".yaml",
    ".xml",
    "_",
    "配置",
    "模板",
    "定义",
    "路由表",
    "接口",
    "日志格式",
    "报文",
)
_MERGE_STRUCTURED_SUFFIXES = (
    "故障类型",
    "消息类型",
    "路由策略",
    "监测点",
    "生命周期",
    "设计哲学",
)
_MERGE_TITLE_WRAPPER_MARKERS = (
    "管理",
    "操作",
    "完整",
    "当前",
    "默认",
    "整体",
    "全局",
    "全系统",
    "项目",
    "年度",
    "季度",
    "月度",
    "报告",
    "调查",
    "复盘",
    "方案",
    "计划",
    "记录",
    "日志",
    "流程",
    "步骤",
    "调试",
    "执行",
    "部署",
    "优化",
    "行动",
    "技术",
    "系统",
    "运行",
    "支持",
    "checklist",
    "清单",
)
_BAD_TERM_PREFIXES = ("在", "由", "将", "每", "各", "该", "此", "其", "所有", "全部", "相关", "主要")
_BAD_TERM_SUFFIXES = ("是", "的", "了", "内", "外", "时", "中", "后", "前", "于", "及", "和")
_SECTION_TERM_SUFFIXES = (
    "控制模块",
    "工具模块",
    "实现模块",
    "通信模块",
    "处理器",
    "流程",
    "恢复",
    "准备",
    "策略",
    "规范",
    "建议",
    "对抗",
    "分析",
    "现状",
    "管理",
    "窗口",
    "要点",
    "内容",
    "方法",
    "工具",
    "优化",
    "总结",
    "处置",
    "处理",
    "验证",
    "总览",
    "统计",
    "记录",
    "达标率",
)
_CAUSAL_MARKERS = ("导致", "根因", "误判", "漏检", "缺陷", "故障", "异常", "触发", "延迟", "安排", "换羽", "隔离")
_GENERIC_SUMMARY_MARKERS = ("核心纪律", "核心能力", "一种模式", "一种状态", "关键概念", "相关系统")
_SUMMARY_BOOST_MARKERS = (
    "核心资源",
    "物质基础",
    "管理系统",
    "控制中枢",
    "异常现象",
    "初始化仪式",
    "首次受控放量",
    "缓冲节点",
    "中继节点",
    "负载均衡",
    "监测死角",
    "有害副产物",
    "健康标志",
    "光泽",
)
_HIGH_SIGNAL_SCOPE_MARKERS = (
    "核心资源",
    "物质基础",
    "管理系统",
    "自动化控制",
    "调度",
    "联动",
    "触发",
    "阈值",
    "反射率",
    "晦暗",
    "溢彩",
    "负载均衡",
    "中继",
    "缓冲",
    "偏远区域",
    "种月",
    "盲区",
    "覆盖",
    "漏检",
    "底噪",
    "毛刺",
    "峰谷差",
    "幽灵读数",
    "留声机",
    "账房",
    "收支不平",
    "暗流",
    "步骤",
    "阶段",
    "条件",
    "范围",
    "区间",
    "比例",
    "历史记录",
)
_STRUCTURAL_SUPPORT_SUBJECT_PATTERNS = (
    re.compile(r"(?:由|通过|依靠|借助)([\u4e00-\u9fff]{2,8})(?:承担|执行|完成|实现|负责)"),
    re.compile(r"([\u4e00-\u9fff]{2,8})节点(?:在|的|可|需|应|用于|作为|承担)"),
    re.compile(r"([\u4e00-\u9fff]{2,8})技术(?:在|的|可|需|应|用于|作为|让|能|不能|只能)"),
    re.compile(r"^([\u4e00-\u9fff]{2,8})数据质量"),
)
_DOC_SUBJECT_SUFFIXES = (
    "事故复盘报告",
    "事故复盘",
    "应急预案",
    "操作Checklist",
    "操作checklist",
    "Checklist",
    "checklist",
    "操作手册",
    "用户手册",
    "技术白皮书",
    "实施指南",
    "内部培训",
    "培训大纲",
    "应对策略",
    "最佳实践",
    "经验总结",
    "架构设计",
    "设计规范",
    "风险评估报告",
    "管理办法",
    "行动守则",
    "工作总结",
    "年度工作总结",
    "月度趋势报告",
    "周报模板",
    "操作注意事项",
    "部署位置优化",
    "生命周期管理",
    "等级评定标准",
    "分析报告",
    "接口文档",
    "进度汇报",
    "调优记录",
    "评估报告",
    "治理文件",
    "任务调度配置",
    "审计系统",
)
_DOC_SUBJECT_TRIM_SUFFIXES = (
    "流程",
    "操作",
    "步骤",
    "预案",
    "指南",
    "策略",
    "报告",
    "系统",
    "协议",
)
_ARTIFACT_SUBJECT_REWRITES = (
    ("路由表", "路由策略"),
    ("报文定义", "消息类型"),
    ("监测点配置", "监测点"),
)
_DOC_SECTION_MARKERS = (
    "定义",
    "前置条件",
    "启动条件",
    "分级",
    "职责",
    "流程",
    "步骤",
    "执行",
    "检查",
    "验证",
    "准备",
    "处理",
    "处置",
    "回退",
    "解除",
    "要求",
    "根因",
    "原因",
    "影响",
    "改进",
    "总结",
    "特征",
    "运行模式",
    "控制点",
    "时间线",
    "总览",
    "统计",
    "记录",
    "达标率",
    "质量",
)
_STRUCTURAL_SCOPE_MARKERS = (
    "阈值",
    "触发",
    "周期",
    "范围",
    "数量",
    "类型",
    "流程",
    "步骤",
    "阶段",
    "稳定性",
    "覆盖率",
    "达标率",
    "底噪",
    "峰谷差",
    "毛刺",
    "盲区",
    "部署",
    "中继",
    "缓冲",
    "策略",
    "计划",
    "条件",
)
_STRUCTURAL_SUBJECT_SUFFIXES = (
    "数据质量",
    "历史记录完整性",
    "历史记录",
    "部署策略",
    "路由策略",
    "路由表",
    "生命周期",
    "触发阈值",
    "触发条件",
    "任务调度配置",
    "自动化控制系统",
    "审计系统",
    "周报模板",
)
_MERGEABLE_TITLE_SUFFIXES = (
    "流程要点",
    "处理流程",
    "流程规范",
    "操作流程",
    "应急预案",
    "预警机制",
    "预警与应对",
    "风险评估报告",
    "风险评估",
    "协作图",
    "活动图",
    "状态机",
    "架构图",
    "组件图",
    "周报模板",
    "任务调度配置",
    "调度配置",
    "审计系统",
    "材质参数",
    "自动化控制系统",
    "自动化控制",
    "调试模式",
    "管理",
    "执行",
    "后验证",
    "记录",
    "团队",
)


@dataclass
class ConceptDraft:
    name: str
    aliases: set[str] = field(default_factory=set)
    sources: set[str] = field(default_factory=set)
    related: set[str] = field(default_factory=set)
    categories: set[str] = field(default_factory=set)
    summary_candidates: list[tuple[int, str]] = field(default_factory=list)
    scope_parts: list[str] = field(default_factory=list)

    def add_summary(self, text: str, *, priority: int) -> None:
        normalized = _strip_contrastive_tail(_normalize_sentence(text))
        if not normalized:
            return
        self.summary_candidates.append((priority, normalized))

    def add_scope(self, text: str) -> None:
        normalized = _normalize_sentence(text)
        if not normalized:
            return
        if normalized not in self.scope_parts:
            self.scope_parts.append(normalized)

    def best_summary(self) -> str:
        if not self.summary_candidates:
            return ""
        ordered = sorted(
            self.summary_candidates,
            key=lambda item: (
                -(item[0] + _summary_quality_bonus(item[1], term=self.name)),
                len(item[1]),
            ),
        )
        return ordered[0][1]


@dataclass
class SourceRecord:
    path: Path
    source_name: str
    text: str
    sentences: list[str]


class OfflineKnowledgeBuilder:
    def __init__(self, kb_root: str | Path) -> None:
        self.kb_root = Path(kb_root)
        self.entries_dir = self.kb_root / "entries"
        self.placeholders_dir = self.kb_root / "placeholders"
        self.indexes_dir = self.kb_root / "indexes"
        self.drafts: dict[str, ConceptDraft] = {}
        self.records: list[SourceRecord] = []

    def load_existing_kb(self) -> None:
        if not self.kb_root.exists():
            return
        try:
            data = kb_inventory(self.kb_root)
        except Exception:
            return
        for name, doc in data.get("docs", {}).items():
            if doc.get("kind") != "formal":
                continue
            draft = self._draft(name)
            draft.sources.update(str(item).strip() for item in doc.get("sources", []) if str(item).strip())
            draft.aliases.update(str(item).strip() for item in doc.get("aliases", []) if str(item).strip())
            summary = str(doc.get("summary", "")).strip()
            if summary:
                draft.add_summary(summary, priority=_loaded_summary_priority(name, summary))
            for section_name, content in (doc.get("sections_map", {}) or {}).items():
                if section_name == "Related":
                    for target in re.findall(r"\[\[([^\]]+)\]\]", str(content)):
                        cleaned = target.split("|")[0].split("#")[0].strip()
                        if cleaned and cleaned != name:
                            draft.related.add(cleaned)
                    continue
                draft.add_scope(str(content))

    def ingest_materials(self, materials: Iterable[str | Path]) -> dict[str, Any]:
        for raw_path in materials:
            path = Path(raw_path)
            if not path.exists() or not path.is_file():
                continue
            record = _load_source_record(path)
            self.records.append(record)
            self._ingest_record(record)
        self._ingest_frequent_terms()
        self._enrich_from_records()
        self._promote_canonical_bare_terms()
        self._merge_supporting_title_variants()
        self._fanout_structural_scope_lines()
        self._stabilize_summaries()
        self._write_entries()
        self._write_indexes()
        return {
            "entry_count": len(list(self.entries_dir.glob("*.md"))),
            "placeholder_count": len(list(self.placeholders_dir.glob("*.md"))),
        }

    def tidy(self, *, focus: str = "general") -> dict[str, Any]:
        _ = focus
        self.load_existing_kb()
        self._enrich_from_records()
        self._promote_canonical_bare_terms()
        self._merge_supporting_title_variants()
        self._fanout_structural_scope_lines()
        self._stabilize_summaries()
        self._write_entries()
        self._write_indexes()
        return {
            "entry_count": len(list(self.entries_dir.glob("*.md"))),
            "placeholder_count": len(list(self.placeholders_dir.glob("*.md"))),
        }

    def _ingest_record(self, record: SourceRecord) -> None:
        suffix = record.path.suffix.lower()
        if suffix == ".md":
            self._ingest_markdown(record)
        elif suffix in {".py"}:
            self._ingest_python_code(record)
        elif suffix in {".cpp", ".h"}:
            self._ingest_cpp_code(record)
        elif suffix in {".yaml", ".yml"}:
            self._ingest_yaml(record)
        elif suffix == ".json":
            self._ingest_json(record)
        elif suffix == ".xml":
            self._ingest_xml(record)
        else:
            self._ingest_textual(record)

    def _ingest_markdown(self, record: SourceRecord) -> None:
        self._ingest_markdown_sections(
            record,
            doc_subjects=_document_subject_terms(record),
        )
        self._ingest_markdown_bullets(record)
        self._ingest_textual(record)

    def _ingest_markdown_sections(
        self,
        record: SourceRecord,
        *,
        doc_subjects: list[str] | None = None,
    ) -> None:
        current_heading = ""
        current_body: list[str] = []
        subjects = doc_subjects or []
        for raw_line in record.text.splitlines():
            line = raw_line.rstrip()
            heading_match = re.match(r"^\s*#{1,6}\s+(.+?)\s*$", line)
            if heading_match:
                if current_heading and current_body:
                    self._ingest_markdown_section(
                        record,
                        current_heading,
                        current_body,
                        doc_subjects=subjects,
                    )
                current_heading = heading_match.group(1)
                current_body = []
                continue
            if current_heading:
                current_body.append(line)
        if current_heading and current_body:
            self._ingest_markdown_section(
                record,
                current_heading,
                current_body,
                doc_subjects=subjects,
            )

    def _ingest_markdown_section(
        self,
        record: SourceRecord,
        heading: str,
        body_lines: list[str],
        *,
        doc_subjects: list[str] | None = None,
    ) -> None:
        terms = _extract_heading_terms(heading)
        if not terms:
            terms = []
        body_text = "\n".join(line for line in body_lines if line.strip())
        sentences = [
            _normalize_sentence(chunk)
            for chunk in _SENTENCE_PATTERN.split(body_text)
            if _normalize_sentence(chunk)
        ]
        informative_lines = [
            _normalize_sentence(line)
            for line in body_lines
            if line.strip() and not line.strip().startswith("*相关文档")
        ]
        target_terms = list(terms)
        if doc_subjects and _should_project_section_to_doc_subject(heading, informative_lines):
            for subject in doc_subjects:
                if subject not in target_terms:
                    target_terms.append(subject)

        for term in target_terms:
            draft = self._draft(term)
            draft.sources.add(record.source_name)
            summary = (
                self._find_best_summary(term, sentences)
                or _find_compound_reference_summary(term, sentences)
                or (
                    _best_heading_body_summary(term, sentences)
                    if term in terms or _can_project_section_summary(term, heading, sentences)
                    else ""
                )
            )
            if summary:
                draft.add_summary(summary, priority=165 if term in terms else 156)
            for line in _select_informative_lines(
                informative_lines,
                name=term,
                related=draft.related,
                limit=4 if term in terms else 3,
            ):
                draft.add_scope(line)

    def _ingest_markdown_bullets(self, record: SourceRecord) -> None:
        for raw_line in record.text.splitlines():
            stripped = raw_line.strip()
            if not stripped.startswith(("-", "*", "•")):
                continue
            line = re.sub(r"^[-*•]\s*", "", stripped)
            for term, summary in _extract_bullet_fact_candidates(line):
                draft = self._draft(term)
                draft.sources.add(record.source_name)
                draft.add_summary(summary, priority=162)
                draft.add_scope(summary)

    def _ingest_python_code(self, record: SourceRecord) -> None:
        self._ingest_code_todos(record)
        source_text = textwrap.dedent(record.text)
        try:
            tree = ast.parse(source_text)
        except SyntaxError:
            self._ingest_textual(record)
            return

        module_doc = ast.get_docstring(tree) or ""
        subject = _extract_code_subject(module_doc, fallback=record.path.stem)
        exception_terms: list[str] = []

        for node in tree.body:
            if isinstance(node, ast.ClassDef):
                doc = ast.get_docstring(node) or ""
                if node.name.endswith("Error"):
                    label = _exception_label_from_doc(doc, node.name)
                    if label and not _is_generic_exception_label(label, subject):
                        exception_terms.append(label)
                        draft = self._draft(label)
                        draft.sources.add(record.source_name)
                        draft.categories.add("safety")
                        draft.aliases.add(node.name)
                        if doc and not _is_stub_exception_doc(label, doc):
                            draft.add_summary(_ensure_subject_sentence(label, doc), priority=170)
                else:
                    class_subject = _extract_code_subject(doc, fallback="")
                    if class_subject and class_subject != subject:
                        class_draft = self._draft(class_subject)
                        class_draft.sources.add(record.source_name)
                        class_draft.aliases.add(node.name)
                        if doc:
                            class_draft.add_summary(_ensure_subject_sentence(class_subject, doc), priority=150)
                for child in node.body:
                    if isinstance(child, (ast.FunctionDef, ast.AsyncFunctionDef)):
                        self._ingest_code_callable_doc(
                            record,
                            child.name,
                            ast.get_docstring(child) or "",
                            fallback_subject=subject,
                        )
            elif isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
                self._ingest_code_callable_doc(
                    record,
                    node.name,
                    ast.get_docstring(node) or "",
                    fallback_subject=subject,
                )

        taxonomy_name = _clean_term(f"{subject}故障类型") if subject else ""
        if taxonomy_name and exception_terms:
            taxonomy = self._draft(taxonomy_name)
            taxonomy.sources.add(record.source_name)
            taxonomy.categories.add("safety")
            taxonomy.aliases.update({record.path.name, record.path.stem, "异常类"})
            taxonomy.add_summary(
                f"{taxonomy_name}包括{_join_cn(exception_terms)}等运行时异常。",
                priority=176,
            )
            for term in exception_terms:
                taxonomy.related.add(term)

        if subject:
            subject_draft = self._draft(subject)
            subject_draft.sources.add(record.source_name)
            subject_draft.aliases.add(record.path.stem)
            if module_doc:
                subject_draft.add_summary(_ensure_subject_sentence(subject, module_doc.splitlines()[0]), priority=148)

        self._ingest_textual(record)

    def _ingest_cpp_code(self, record: SourceRecord) -> None:
        self._ingest_code_todos(record)
        subject = _extract_code_subject(record.text, fallback=record.path.stem)
        exception_terms: list[str] = []
        for class_name, doc in re.findall(
            r"class\s+([A-Za-z0-9_]+Error)[^{]*\{.*?@brief\s+([^\n*]+)",
            record.text,
            flags=re.DOTALL,
        ):
            label = _exception_label_from_doc(doc, class_name)
            if not label or _is_generic_exception_label(label, subject):
                continue
            exception_terms.append(label)
            draft = self._draft(label)
            draft.sources.add(record.source_name)
            draft.categories.add("safety")
            draft.aliases.add(class_name)
            if not _is_stub_exception_doc(label, doc):
                draft.add_summary(_ensure_subject_sentence(label, doc), priority=168)

        for brief, method_name in re.findall(
            r"@brief\s+([^\n*]+).*?\b([A-Za-z_][A-Za-z0-9_]*)\s*\(",
            record.text,
            flags=re.DOTALL,
        ):
            self._ingest_code_callable_doc(
                record,
                method_name,
                brief.strip(),
                fallback_subject=subject,
            )

        taxonomy_name = _clean_term(f"{subject}故障类型") if subject else ""
        if taxonomy_name and exception_terms:
            taxonomy = self._draft(taxonomy_name)
            taxonomy.sources.add(record.source_name)
            taxonomy.categories.add("safety")
            taxonomy.aliases.update({record.path.name, record.path.stem, "异常类"})
            taxonomy.add_summary(
                f"{taxonomy_name}包括{_join_cn(exception_terms)}等运行时异常。",
                priority=176,
            )
            for term in exception_terms:
                taxonomy.related.add(term)

        self._ingest_textual(record)

    def _ingest_code_callable_doc(
        self,
        record: SourceRecord,
        callable_name: str,
        docstring: str,
        *,
        fallback_subject: str = "",
    ) -> None:
        if not docstring:
            return
        normalized = "\n".join(line.strip() for line in docstring.splitlines() if line.strip())
        first_line = normalized.splitlines()[0]
        explicit_title, description = _split_code_doc_title(first_line)
        target_name = explicit_title or _extract_code_subject(first_line, fallback=fallback_subject)
        if not target_name:
            return
        draft = self._draft(target_name)
        draft.sources.add(record.source_name)
        if callable_name:
            draft.aliases.add(callable_name)
        summary = description or normalized
        candidate_summary = _ensure_subject_sentence(target_name, summary)
        if (
            not explicit_title
            and target_name == fallback_subject
            and target_name not in normalized
            and _summary_quality_bonus(candidate_summary, term=target_name) < 0
        ):
            if "目标" in normalized or "共振峰" in normalized or "重试" in normalized:
                draft.add_scope(_normalize_sentence(normalized))
            return
        draft.add_summary(candidate_summary, priority=166)
        if "目标" in normalized or "共振峰" in normalized or "重试" in normalized:
            draft.add_scope(_normalize_sentence(normalized))

    def _ingest_code_todos(self, record: SourceRecord) -> None:
        todos = [
            _normalize_sentence(match)
            for match in re.findall(r"TODO[:：]?\s*([^\n]+)", record.text)
        ]
        todos = [item for item in todos if item]
        if not todos:
            return
        draft = self._draft("代码待实现项")
        draft.sources.add(record.source_name)
        draft.categories.update({"operations", "core"})
        draft.aliases.update({"TODO", "待实现功能", "未完成功能"})
        categories = []
        for item in todos:
            for marker in ("迁移", "检测", "通知", "审计", "判官", "账房", "留声机", "信使"):
                if marker in item and marker not in categories:
                    categories.append(marker)
        if categories:
            draft.add_summary(
                f"代码待实现项记录了仍需接入真实环境的{_join_cn(categories[:5])}集成功能。",
                priority=178,
            )
        else:
            draft.add_summary("代码待实现项记录了当前仍未落地的真实环境集成功能。", priority=178)
        for item in todos[:8]:
            draft.add_scope(item)

    def _ingest_yaml(self, record: SourceRecord) -> None:
        try:
            payload = yaml.safe_load(record.text)
        except yaml.YAMLError:
            self._ingest_textual(record)
            return
        if not isinstance(payload, dict):
            self._ingest_textual(record)
            return

        if "roles" in payload and isinstance(payload["roles"], dict):
            self._ingest_role_config(record, payload)
        if "permissions" in payload and isinstance(payload["permissions"], dict):
            self._ingest_permission_config(record, payload)
        if any(key in payload for key in ("monitoring", "resonator", "security", "design_philosophy")):
            self._ingest_system_config(record, payload)
        if "rules" in payload and isinstance(payload["rules"], list):
            self._ingest_alert_rules(record, payload)

        self._ingest_textual(record)

    def _ingest_role_config(self, record: SourceRecord, payload: dict[str, Any]) -> None:
        for role_key, role in payload.get("roles", {}).items():
            if not isinstance(role, dict):
                continue
            name = _clean_term(role.get("display_name"))
            if not name:
                continue
            draft = self._draft(name)
            draft.sources.add(record.source_name)
            draft.categories.add("roles")
            draft.aliases.update(_string_set(role_key, role.get("category")))
            description = str(role.get("description", "")).strip()
            if description:
                draft.add_summary(_ensure_subject_sentence(name, description), priority=182)
                draft.add_scope(description)
            if role.get("read_only"):
                draft.add_scope(f"{name}默认只有只读观察权限，不能执行生产操作。")
            if role.get("promotion_target"):
                draft.add_scope(f"{name}的默认晋升目标是{role['promotion_target']}对应的正式值班角色。")

    def _ingest_permission_config(self, record: SourceRecord, payload: dict[str, Any]) -> None:
        for domain_name, domain in payload.get("permissions", {}).items():
            if not isinstance(domain, dict):
                continue
            for permission_key, permission in domain.items():
                if not isinstance(permission, dict):
                    continue
                name = _clean_term(permission.get("name"))
                if not name:
                    continue
                draft = self._draft(name)
                draft.sources.add(record.source_name)
                draft.categories.add(_category_for_domain(domain_name))
                draft.aliases.update(_string_set(permission_key))
                description = str(permission.get("description", "")).strip()
                if description:
                    draft.add_summary(_ensure_subject_sentence(name, description), priority=176)
                    draft.add_scope(description)
                allowed_roles = [
                    str(item).strip()
                    for item in permission.get("allowed_roles", [])
                    if str(item).strip()
                ]
                if allowed_roles:
                    role_names = _resolve_role_names(payload.get("roles", {}), allowed_roles)
                    if role_names:
                        draft.add_scope(f"{name}的默认授权角色包括{_join_cn(role_names)}。")
                        draft.related.update(role_names)

    def _ingest_system_config(self, record: SourceRecord, payload: dict[str, Any]) -> None:
        monitoring = payload.get("monitoring") or {}
        if isinstance(monitoring, dict):
            hum_level = monitoring.get("hum_level") or {}
            if isinstance(hum_level, dict):
                hum_level_draft = self._draft("嗡鸣度")
                hum_level_draft.sources.add(record.source_name)
                hum_level_draft.categories.add("monitoring")
                hum_level_draft.aliases.update({"hum_level", "嗡鸣度指标"})
                hum_level_draft.add_summary(
                    "嗡鸣度是衡量哈基米活跃状态和系统整体能量活跃程度的核心监控指标。",
                    priority=188,
                )
                sample_rate = hum_level.get("sample_rate")
                if sample_rate is not None:
                    hum_level_draft.add_scope(f"嗡鸣度默认以 {sample_rate}Hz 的采样频率持续监测。")
                refresh = hum_level.get("echo_wall_refresh")
                if refresh is not None:
                    hum_level_draft.add_scope(f"回音壁的嗡鸣度刷新间隔默认为 {refresh} 秒。")

                resonance_peak = hum_level.get("resonance_peak") or {}
                lower = resonance_peak.get("min")
                upper = resonance_peak.get("max")
                if lower is not None and upper is not None:
                    peak = self._draft("共振峰")
                    peak.sources.add(record.source_name)
                    peak.categories.add("monitoring")
                    peak.aliases.update({"resonance_peak", "共振峰区间"})
                    peak.add_summary(
                        f"共振峰是嗡鸣度的理想运行区间，标准腔通常保持在 {lower}-{upper}Hz。",
                        priority=190,
                    )
                    peak.add_scope("超出共振峰会增加毛刺和分流处置风险。")
                    hum_level_draft.related.add("共振峰")
                    hum_level_draft.add_scope(f"嗡鸣度的标准腔共振峰范围为 {lower}-{upper}Hz。")

                red_line = hum_level.get("red_line")
                if red_line is not None:
                    red_line_draft = self._draft("红线")
                    red_line_draft.sources.add(record.source_name)
                    red_line_draft.categories.update({"monitoring", "safety"})
                    red_line_draft.aliases.add("red_line")
                    red_line_draft.add_summary(
                        f"红线是嗡鸣度绝对不可超过的安全阈值，标准腔默认值为 {red_line}Hz。",
                        priority=189,
                    )
                    hum_level_draft.related.add("红线")
                    hum_level_draft.add_scope(f"嗡鸣度触及 {red_line}Hz 红线时必须立即进入保护或处置流程。")

                base_noise = hum_level.get("base_noise")
                if base_noise is not None:
                    base_noise_draft = self._draft("底噪")
                    base_noise_draft.sources.add(record.source_name)
                    base_noise_draft.categories.add("monitoring")
                    base_noise_draft.aliases.update({"base_noise", "底噪声"})
                    base_noise_draft.add_summary(
                        f"底噪是系统固有最低嗡鸣度的基线读数，标准腔默认值为 {base_noise}Hz。",
                        priority=186,
                    )
                    hum_level_draft.related.add("底噪")

            safety_string = str(monitoring.get("safety_string", "")).strip()
            if safety_string:
                safety = self._draft("安全弦")
                safety.sources.add(record.source_name)
                safety.categories.update({"monitoring", "safety"})
                safety.aliases.add("safety_string")
                safety.add_summary("安全弦是嗡鸣度位于底噪以上、共振峰内且红线以下的稳定运行区间。", priority=187)
                safety.add_scope(f"系统配置中的安全弦约束为 {safety_string}。")

        resonator = payload.get("resonator") or {}
        if isinstance(resonator, dict):
            purity_ratio = resonator.get("purity_ratio") or {}
            if isinstance(purity_ratio, dict):
                draft = self._draft("清浊比")
                draft.sources.add(record.source_name)
                draft.categories.update({"monitoring", "core"})
                draft.aliases.update({"purity_ratio", "ratio"})
                target = purity_ratio.get("target")
                warning = purity_ratio.get("warning")
                emergency = purity_ratio.get("emergency")
                draft.add_summary(
                    "清浊比是纯净哈基米与散斑的比例指标，用于衡量谐振腔和传输链路的健康度。",
                    priority=186,
                )
                if target is not None:
                    draft.add_scope(f"日常运维的目标清浊比为 {target}（约 95:5）。")
                if warning is not None:
                    draft.add_scope(f"清浊比低于 {warning}（约 90:10）时应启动清道夫流程。")
                if emergency is not None:
                    draft.add_scope(f"清浊比低于 {emergency}（约 85:15）时应暂停腔体运行并准备补天。")

        security = payload.get("security") or {}
        if isinstance(security, dict):
            rule = security.get("three_strike_rule") or {}
            if isinstance(rule, dict) and rule.get("enabled"):
                window = rule.get("rolling_window_hours")
                duration = rule.get("isolation_duration_hours")
                draft = self._draft("三振法则")
                draft.sources.add(record.source_name)
                draft.categories.add("safety")
                draft.aliases.update({"three_strike_rule", "连续三次异常自动隔离"})
                summary = "三振法则是在滚动窗口内连续三次出现同类异常时自动隔离故障单元的安全规则。"
                if window is not None:
                    summary = f"三振法则是在 {window} 小时滚动窗口内连续三次出现同类异常时自动隔离故障单元的安全规则。"
                draft.add_summary(summary, priority=190)
                if duration is not None:
                    draft.add_scope(f"触发三振法则后的默认隔离时长为 {duration} 小时。")
                if rule.get("require_manual_review"):
                    draft.add_scope("三振法则触发后仍需人工复核，避免把误报直接固化为最终处置。")

        philosophy = payload.get("design_philosophy") or {}
        if isinstance(philosophy, dict):
            self._ingest_design_philosophy(record, philosophy)

    def _ingest_design_philosophy(self, record: SourceRecord, philosophy: dict[str, Any]) -> None:
        draft = self._draft("哈基米系统设计哲学")
        draft.sources.add(record.source_name)
        draft.categories.update({"core", "operations"})
        draft.aliases.update({"设计哲学", "system_design_philosophy"})
        principles: list[str] = []
        for key, item in philosophy.items():
            if not isinstance(item, dict):
                continue
            description = _normalize_sentence(item.get("description"))
            if description:
                principles.append(description.rstrip("。"))
                draft.add_scope(description)
            mechanisms = [
                _normalize_sentence(value)
                for value in item.get("mechanisms", [])
                if _normalize_sentence(value)
            ]
            for mechanism in mechanisms[:4]:
                draft.add_scope(mechanism)
                for target in _extract_definition_subjects(mechanism):
                    draft.related.add(target)
        if principles:
            draft.add_summary(
                f"哈基米系统设计哲学强调{_join_cn(principles[:5])}。",
                priority=184,
            )
        system_draft = self._draft("哈基米系统")
        system_draft.sources.add(record.source_name)
        system_draft.related.add("哈基米系统设计哲学")

    def _ingest_alert_rules(self, record: SourceRecord, payload: dict[str, Any]) -> None:
        rules = [item for item in payload.get("rules", []) if isinstance(item, dict)]
        config_name = "告警规则配置"
        draft = self._draft(config_name)
        draft.sources.add(record.source_name)
        draft.categories.update({"monitoring", "safety"})
        draft.aliases.update({record.path.name, record.path.stem, "告警规则集"})
        rules_count = payload.get("global", {}).get("rules_count")
        if rules_count:
            draft.add_summary(
                f"{config_name}是维护系统阈值、升级链路和应急动作的集中规则集，当前定义 {rules_count} 条规则。",
                priority=135,
            )
            draft.add_scope(f"当前全局规则数量为 {rules_count} 条。")
        severities = payload.get("global", {}).get("severity_levels") or []
        if severities:
            draft.add_scope(f"支持的告警级别包括{_join_cn(str(item) for item in severities)}。")
        escalation = payload.get("global", {}).get("escalation", {}).get("levels", [])
        escalation_targets = []
        for item in escalation:
            if isinstance(item, dict) and item.get("target"):
                escalation_targets.append(str(item["target"]).strip())
        if escalation_targets:
            draft.add_scope(f"升级链路会依次通知{_join_cn(escalation_targets)}。")
        for rule in rules[:8]:
            category = _clean_term(rule.get("category"))
            if category:
                draft.related.add(category)

    def _ingest_json(self, record: SourceRecord) -> None:
        try:
            payload = json.loads(record.text)
        except json.JSONDecodeError:
            self._ingest_textual(record)
            return
        if not isinstance(payload, dict):
            self._ingest_textual(record)
            return

        metrics = payload.get("metrics")
        if isinstance(metrics, list):
            self._ingest_metrics(record, metrics)
        if isinstance(payload.get("nodes"), list):
            self._ingest_topology(record, payload)

        self._ingest_textual(record)

    def _ingest_metrics(self, record: SourceRecord, metrics: list[Any]) -> None:
        for metric in metrics:
            if not isinstance(metric, dict):
                continue
            name = _clean_term(metric.get("name"))
            if not name:
                continue
            draft = self._draft(name)
            draft.sources.add(record.source_name)
            draft.categories.add("monitoring")
            draft.aliases.update(_string_set(metric.get("code"), *(metric.get("unit_aliases") or [])))
            description = str(metric.get("description", "")).strip()
            if description:
                draft.add_summary(_ensure_subject_sentence(name, description), priority=145)
            structured_summary = _metric_structured_summary(name, metric)
            if structured_summary:
                draft.add_summary(structured_summary, priority=172)
            unit = str(metric.get("unit", "")).strip()
            if unit:
                unit_aliases = [str(item).strip() for item in metric.get("unit_aliases", []) if str(item).strip()]
                display_units = [unit, *unit_aliases]
                draft.add_scope(f"{name}默认以{_join_cn(display_units)}表示。")
            thresholds = metric.get("thresholds") or {}
            threshold_bits = []
            for key, value in thresholds.items():
                if not isinstance(value, dict):
                    continue
                desc = str(value.get("description", "")).strip()
                number = value.get("value")
                if desc:
                    threshold_bits.append(desc)
                elif number is not None:
                    threshold_bits.append(f"{key}阈值为 {number}")
            for bit in threshold_bits[:4]:
                draft.add_scope(bit)
            related_issues = [str(item).strip() for item in metric.get("related_issues", []) if str(item).strip()]
            draft.related.update(related_issues)

    def _ingest_topology(self, record: SourceRecord, payload: dict[str, Any]) -> None:
        topology_summary = payload.get("topology_summary", {}) or {}
        summary = str(topology_summary.get("summary_text", "")).strip()
        if summary:
            draft = self._draft("部署拓扑")
            draft.sources.add(record.source_name)
            draft.categories.update({"network", "core"})
            draft.aliases.update({record.path.name, record.path.stem})
            draft.add_summary(_ensure_subject_sentence("部署拓扑", summary), priority=110)
            station_strategy = str(topology_summary.get("station_strategy", "")).strip()
            if station_strategy:
                draft.add_scope(station_strategy)
                station = self._draft("驿站")
                station.sources.add(record.source_name)
                station.categories.add("network")
                station.aliases.add("驿站节点")
                station.add_scope(station_strategy)
                station.add_summary(
                    "驿站是部署在高频路径、跨区边界和偏远接入区的中继与缓冲节点，用于负载均衡与临时缓存。",
                    priority=188,
                )
        resonator_roles = topology_summary.get("resonator_roles") or {}
        hot_backup_count = resonator_roles.get("hot_backup")
        if isinstance(hot_backup_count, int) and hot_backup_count > 0:
            draft = self._draft("热备份")
            draft.sources.add(record.source_name)
            draft.categories.update({"core", "operations"})
            draft.aliases.add("热备份谐振腔")
            draft.add_summary(
                f"热备份是用于在主谐振腔失效时立即接管负载的备用谐振腔角色，当前拓扑包含 {hot_backup_count} 个热备份节点。",
                priority=164,
            )
            draft.add_scope("热备份节点保持随时可接管状态，通常与主谐振腔成对部署。")
        for node in payload.get("nodes", []):
            if not isinstance(node, dict):
                continue
            deployment_mode = _clean_term(node.get("deployment_mode"))
            if not deployment_mode:
                deployment_mode = _clean_term((node.get("config") or {}).get("deployment_mode"))
            if deployment_mode:
                draft = self._draft(deployment_mode)
                draft.sources.add(record.source_name)
                draft.categories.add("network")
                description = f"{deployment_mode}是用于边缘或偏远节点接入的小型部署模式。"
                draft.add_summary(description, priority=105)
                role = str(node.get("role", "")).strip()
                if role:
                    draft.add_scope(f"{deployment_mode}节点通常承担{role}角色。")
            role = str(node.get("role", "")).strip()
            if role == "hot_backup":
                draft = self._draft("热备份")
                draft.sources.add(record.source_name)
                draft.categories.update({"core", "operations"})
                draft.add_scope(f"热备份节点通常部署在{str(node.get('zone', '')).strip() or '关键区域'}，用于主腔故障接管。")

    def _ingest_xml(self, record: SourceRecord) -> None:
        if "<enum" in record.text and "消息类型" in record.text:
            labels = re.findall(r'<enum[^>]+label="([^"]+)"', record.text)
            draft = self._draft("旋涡协议消息类型")
            draft.sources.add(record.source_name)
            draft.categories.update({"transfer", "core"})
            draft.aliases.update({record.path.name, record.path.stem})
            draft.add_summary("旋涡协议消息类型定义了标准传输、跃迁、晶格化和剥离等合法报文类别。", priority=184)
            if labels:
                draft.add_scope(f"当前消息类型包括{_join_cn(labels)}。")
            draft.related.update({"旋涡协议", "跃迁", "晶格化", "剥离"})
        if "<task " in record.text:
            self._ingest_xml_tasks(record)
        sections = re.findall(
            r'<section name="([^"]+)">(.*?)</section>',
            record.text,
            flags=re.DOTALL,
        )
        for section_name, body in sections:
            cleaned_section = _clean_term(section_name)
            if cleaned_section in {"消息类型", "异常状态码", "路由策略", "路由规则", "跃迁规则", "监测点列表", "盲区覆盖策略"}:
                self._ingest_xml_section(record, cleaned_section, body)
        self._ingest_xml_structured_params(record)
        self._ingest_textual(record)

    def _ingest_xml_tasks(self, record: SourceRecord) -> None:
        tasks = re.findall(r'<task\b([^>]*)>(.*?)</task>', record.text, flags=re.DOTALL)
        if not tasks:
            return

        config_draft = self._draft("千机匣任务调度配置")
        config_draft.sources.add(record.source_name)
        config_draft.categories.update({"operations", "core"})
        config_draft.aliases.update({record.path.name, record.path.stem})

        schedule_bits: list[str] = []
        task_names: list[str] = []

        for attrs, body in tasks:
            name_match = re.search(r'name="([^"]+)"', attrs)
            if not name_match:
                continue
            name = _clean_term(name_match.group(1))
            if not name:
                continue
            task_type = str(re.search(r'type="([^"]+)"', attrs).group(1) if re.search(r'type="([^"]+)"', attrs) else "").strip()
            schedule = str(re.search(r'schedule="([^"]+)"', attrs).group(1) if re.search(r'schedule="([^"]+)"', attrs) else "").strip()
            schedule_desc = _describe_schedule(schedule)
            description_match = re.search(r"<description>(.*?)</description>", body, flags=re.DOTALL)
            description = _normalize_sentence(description_match.group(1)) if description_match else ""

            draft = self._draft(name)
            draft.sources.add(record.source_name)
            draft.categories.update({"operations", "core"})
            summary = f"{name}是{schedule_desc}执行的{task_type or '自动化'}任务。"
            if description:
                summary = f"{summary.rstrip('。')} {description}"
            draft.add_summary(summary, priority=168)
            draft.add_scope(f"{name}的执行周期为{schedule_desc}。")
            if task_type:
                draft.add_scope(f"{name}属于{task_type}任务。")
            steps = re.findall(r'action="([^"]+)"', body)
            if steps:
                draft.add_scope(f"{name}的关键步骤包括{_join_cn(steps[:4])}。")
            params = re.findall(
                r'<param name="([^"]+)" value="([^"]+)"(?: unit="([^"]+)")?',
                body,
            )
            for param_name, value, unit in params[:4]:
                unit_suffix = f"{unit}" if unit else ""
                draft.add_scope(f"{name}的{param_name}为{value}{unit_suffix}。")

            task_names.append(name)
            schedule_bits.append(f"{name}{schedule_desc}")

        if schedule_bits:
            config_draft.add_summary(
                f"千机匣任务调度配置定义了{_join_cn(schedule_bits[:3])}等任务的执行周期、步骤和超时参数。",
                priority=180,
            )
        config_draft.related.update(task_names[:8])

    def _ingest_xml_structured_params(self, record: SourceRecord) -> None:
        for param in _extract_xml_params(record.text):
            subject = _structured_param_subject(param["name"])
            if not subject:
                continue
            draft = self._draft(subject)
            draft.sources.add(record.source_name)
            fact_line = _structured_param_fact(
                subject=subject,
                field_name=param["name"],
                value=param["value"],
                unit=param["unit"],
                description=param["description"],
            )
            if fact_line:
                draft.add_scope(fact_line)
                summary = _structural_scope_summary(subject, fact_line)
                if summary:
                    draft.add_summary(summary, priority=181)
                elif (
                    "部署策略" in param["name"]
                    or param["name"].endswith(("触发阈值", "触发条件"))
                    or _has_definition_signal(fact_line)
                ):
                    draft.add_summary(fact_line, priority=178)
            if param["description"]:
                draft.add_scope(param["description"])

        for rule in _extract_xml_rule_facts(record.text):
            action = rule["action"]
            subject = _clean_term(action)
            if not subject:
                continue
            draft = self._draft(subject)
            draft.sources.add(record.source_name)
            fact_line = _normalize_sentence(rule["description"]) or _normalize_sentence(
                f"当{rule['field']}{rule['operator']}{rule['threshold']}时执行{subject}。"
            )
            if fact_line:
                draft.add_scope(fact_line)
                summary = _structural_scope_summary(subject, fact_line)
                if summary:
                    draft.add_summary(summary, priority=180)

    def _ingest_xml_section(self, record: SourceRecord, section_name: str, body: str) -> None:
        if section_name == "消息类型":
            draft = self._draft("旋涡协议消息类型")
            draft.sources.add(record.source_name)
            draft.categories.update({"transfer", "core"})
            draft.aliases.update({record.path.name, record.path.stem})
            labels = re.findall(r'label="([^"]+)"', body)
            draft.add_summary("旋涡协议消息类型定义了标准传输、跃迁、晶格化和剥离等合法报文类别。", priority=184)
            if labels:
                draft.add_scope(f"当前消息类型包括{_join_cn(labels)}。")
            draft.related.update({"旋涡协议", "跃迁", "晶格化", "剥离"})
            return
        if section_name == "监测点列表":
            draft = self._draft("回音壁监测点")
            draft.sources.add(record.source_name)
            draft.categories.update({"monitoring", "core"})
            draft.aliases.update({record.path.name, record.path.stem})
            monitor_types = re.findall(r'type="([^"]+)"', body)
            locations = re.findall(r'location="([^"]+)"', body)
            draft.add_summary(
                "回音壁监测点定义了固定式、嵌入式和移动式回音壁在关键区域的部署位置与采样参数。",
                priority=184,
            )
            if monitor_types:
                draft.add_scope(f"当前监测点类型包括{_join_cn(monitor_types)}。")
            if locations:
                draft.add_scope(f"已覆盖的关键区域包括{_join_cn(locations[:4])}。")
            draft.related.update({"回音壁", "盲区"})
            return
        if section_name == "盲区覆盖策略":
            draft = self._draft("盲区")
            draft.sources.add(record.source_name)
            draft.categories.update({"monitoring", "safety"})
            draft.add_summary(
                "盲区是回音壁无法稳定覆盖的监测死角，需要通过移动式回音壁巡航等补偿策略持续扫描。",
                priority=186,
            )
            coverage_values = re.findall(r'coverage="([^"]+)"', body)
            if coverage_values:
                draft.add_scope(f"当前已识别盲区的覆盖率包括{_join_cn(coverage_values[:4])}。")
            strategy_values = re.findall(r'value="([^"]+)"', body)
            if strategy_values:
                draft.add_scope(f"盲区补偿策略包括{_join_cn(strategy_values[:3])}。")
            draft.related.update({"回音壁", "回音壁监测点"})
            return
        if section_name == "异常状态码":
            labels = re.findall(r'label="([^"]+)"', body)
            for label in labels:
                term = _clean_term(label)
                if not term:
                    continue
                draft = self._draft(term)
                draft.sources.add(record.source_name)
                draft.categories.add("transfer")
                description_match = re.search(
                    rf'label="{re.escape(label)}"\s+description="([^"]+)"',
                    body,
                )
                description = description_match.group(1) if description_match else ""
                if description:
                    draft.add_summary(_ensure_subject_sentence(term, description), priority=132)
                draft.related.add("旋涡协议")
            return
        if section_name in {"路由策略", "路由规则"}:
            draft = self._draft("信使路由策略")
            draft.sources.add(record.source_name)
            draft.categories.update({"network", "transfer"})
            draft.aliases.update({record.path.name, record.path.stem})
            draft.add_summary("信使路由策略定义了走线、接力、分流、合流和跃迁等路径选择与转发方式。", priority=184)
            values = re.findall(r'value="([^"]+)"', body)
            route_types = re.findall(r'<route\b[^>]*type="([^"]+)"', body)
            if route_types:
                draft.add_scope(f"当前路由类型包括{_join_cn(route_types)}。")
                draft.related.update(route_types)
            if values:
                draft.add_scope(f"关键参数包括{_join_cn(value for value in values[:4])}。")
            route_blocks = re.findall(
                r'<route\b([^>]*)>(.*?)</route>',
                body,
                flags=re.DOTALL,
            )
            for attrs, route_body in route_blocks:
                route_type = _clean_term(
                    re.search(r'type="([^"]+)"', attrs).group(1)
                    if re.search(r'type="([^"]+)"', attrs)
                    else ""
                )
                if not route_type:
                    continue
                route_draft = self._draft(route_type)
                route_draft.sources.add(record.source_name)
                route_draft.categories.update({"network", "transfer"})
                route_draft.related.add("信使路由策略")
                route_draft.add_summary(_route_type_summary(route_type), priority=158)
                param_matches = re.findall(
                    r'<param name="([^"]+)" value="([^"]+)"(?: unit="([^"]+)")?',
                    route_body,
                )
                for param_name, value, unit in param_matches[:3]:
                    unit_suffix = f"{unit}" if unit else ""
                    route_draft.add_scope(f"{route_type}的{param_name}为{value}{unit_suffix}。")
                hop_types = re.findall(r'type="([^"]+)"', route_body)
                if hop_types:
                    route_draft.add_scope(f"{route_type}路径会经过{_join_cn(hop_types[:4])}节点。")
            return
        if section_name == "跃迁规则":
            draft = self._draft("跃迁规则")
            draft.sources.add(record.source_name)
            draft.categories.update({"network", "transfer"})
            draft.add_summary("跃迁规则定义了何时启用跃迁通道以及对应的窗口和距离条件。", priority=120)
            thresholds = re.findall(r'threshold="([^"]+)"', body)
            if thresholds:
                draft.add_scope(f"当前规则阈值包括{_join_cn(thresholds)}。")

    def _ingest_textual(self, record: SourceRecord) -> None:
        if "词典" in record.source_name or "核心术语" in record.text:
            self._ingest_glossary_pairs(record)
        doc_subjects: list[str] = []
        if record.path.suffix.lower() != ".md" or any(
            suffix in record.path.stem for suffix in _DOC_SUBJECT_SUFFIXES
        ):
            doc_subjects = _document_subject_terms(record)
        for term in doc_subjects:
            draft = self._draft(term)
            draft.sources.add(record.source_name)
            summary = self._find_best_summary(term, record.sentences) or _find_projected_record_summary(
                term,
                record.sentences,
            )
            if summary:
                draft.add_summary(summary, priority=140)
            for sentence in self._supporting_sentences(term, record.sentences):
                draft.add_scope(sentence)
        for raw_line in record.text.splitlines():
            line = raw_line.strip()
            if not line:
                continue
            for term, summary in _extract_inline_fact_candidates(line):
                draft = self._draft(term)
                draft.sources.add(record.source_name)
                draft.add_summary(summary, priority=172)
                draft.add_scope(summary)
        candidates = self._candidate_terms(record)
        for term in candidates:
            draft = self._draft(term)
            draft.sources.add(record.source_name)
            summary = self._find_best_summary(term, record.sentences) or _find_compound_reference_summary(
                term,
                record.sentences,
            )
            if summary:
                draft.add_summary(summary, priority=95)
            for sentence in self._supporting_sentences(term, record.sentences):
                draft.add_scope(sentence)

    def _ingest_glossary_pairs(self, record: SourceRecord) -> None:
        raw_lines = [
            _strip_structural_prefix(line.strip())
            for line in record.text.splitlines()
            if line.strip()
        ]
        for index, line in enumerate(raw_lines[:-1]):
            name = _clean_term(line)
            if not name:
                continue
            definition = _normalize_sentence(raw_lines[index + 1])
            if not definition or len(definition) < 8:
                continue
            if definition == f"{name}。":
                continue
            if not re.search(r"[\u4e00-\u9fff]", definition):
                continue
            draft = self._draft(name)
            draft.sources.add(record.source_name)
            draft.add_summary(_ensure_subject_sentence(name, definition), priority=155)
            if index + 2 < len(raw_lines):
                maybe_scope = _normalize_sentence(raw_lines[index + 2])
                if maybe_scope and len(maybe_scope) <= 36 and maybe_scope != f"{name}。":
                    draft.add_scope(maybe_scope)

    def _candidate_terms(self, record: SourceRecord) -> list[str]:
        candidates: set[str] = set()
        for sentence in record.sentences:
            stripped = sentence.strip()
            if not stripped:
                continue
            for term in _extract_definition_subjects(stripped):
                candidates.add(term)
            for term in _extract_intro_subjects(stripped):
                candidates.add(term)
            for term in _extract_embedded_compound_subjects(stripped):
                candidates.add(term)
            for term in _extract_structural_subjects(stripped):
                candidates.add(term)
            for term, _kind in _extract_canonical_bare_terms(stripped):
                candidates.add(term)
            heading_match = re.match(r"^(?:#+\s*)?([^\s：:]{2,16})\s*$", stripped)
            if heading_match and _looks_like_concept_name(heading_match.group(1)):
                candidates.add(heading_match.group(1))
            colon_match = re.match(r"^([^\s：:]{2,16})[：:]", stripped)
            if colon_match and _has_definition_signal(stripped):
                candidates.add(colon_match.group(1))
            for term in re.findall(r"\*\*([^*]{2,16})\*\*", stripped):
                candidates.add(term)
        return sorted(
            term
            for term in candidates
            if _clean_term(term)
        )

    def _find_best_summary(self, term: str, sentences: list[str]) -> str:
        best_priority = -1
        best_sentence = ""
        variants = _summary_term_variants(term)
        normalized_variants = {_normalize_summary_surface(item) for item in variants if item}
        for sentence in sentences:
            lowered = sentence.lower()
            normalized_sentence = _normalize_summary_surface(sentence)
            if not any(item.lower() in lowered for item in variants) and not any(
                item and item in normalized_sentence for item in normalized_variants
            ):
                continue
            priority = 0
            for index, pattern in enumerate(_SUMMARY_PATTERNS, start=1):
                if any(pattern.format(term=variant) in sentence for variant in variants):
                    priority = 100 - index
                    break
            if any(
                sentence.startswith(f"{variant}：") or sentence.startswith(f"{variant}:")
                for variant in variants
            ) and _has_definition_signal(sentence):
                priority = max(priority, 92)
            if priority == 0 and any(sentence.startswith(variant) for variant in variants) and _has_direct_definition_signal(sentence):
                priority = 75
            if priority == 0:
                continue
            normalized = _normalize_sentence(sentence)
            if priority > best_priority or (priority == best_priority and len(normalized) > len(best_sentence)):
                best_priority = priority
                best_sentence = normalized
        return best_sentence

    def _supporting_sentences(self, term: str, sentences: list[str]) -> list[str]:
        candidates: list[tuple[tuple[int, int], str]] = []
        seen: set[str] = set()
        variants = _summary_term_variants(term)
        normalized_variants = {_normalize_summary_surface(item) for item in variants if item}
        for sentence in sentences:
            lowered = sentence.lower()
            normalized_sentence = _normalize_summary_surface(sentence)
            if not any(item.lower() in lowered for item in variants) and not any(
                item and item in normalized_sentence for item in normalized_variants
            ):
                continue
            normalized = _normalize_sentence(sentence)
            if not normalized or normalized == term:
                continue
            if any(normalized.startswith(variant) for variant in variants) and any(
                marker.format(term=variant) in normalized
                for marker in _SUMMARY_PATTERNS
                for variant in variants
            ):
                continue
            if normalized in seen:
                continue
            seen.add(normalized)
            score = _line_signal_score(normalized, name=term)
            if any(marker in normalized for marker in _SUMMARY_BOOST_MARKERS):
                score += 3
            if any(marker in normalized for marker in ("不能完全", "只要", "高频路径", "偏远区域", "种月")):
                score += 2
            if len(normalized) >= 28:
                score += 1
            candidates.append(((score, -len(normalized)), normalized))
        candidates.sort(key=lambda item: item[0], reverse=True)
        return [item for _score, item in candidates[:6]]

    def _enrich_from_records(self) -> None:
        known_terms = sorted(self.drafts)
        known_set = set(known_terms)
        for record in self.records:
            for sentence in record.sentences:
                sentence_terms = [
                    term
                    for term in known_terms
                    if term in sentence
                ]
                if len(sentence_terms) < 2:
                    continue
                for term in sentence_terms:
                    draft = self._draft(term)
                    for other in sentence_terms:
                        if other != term and other in known_set:
                            draft.related.add(other)
                    if _has_structural_signal(sentence):
                        draft.add_scope(sentence)

    def _ingest_frequent_terms(self) -> None:
        counts: Counter[str] = Counter()
        all_sentences: list[str] = []
        sources_by_term: dict[str, set[str]] = {}
        for record in self.records:
            for sentence in record.sentences:
                all_sentences.append(sentence)
                local_terms = {
                    cleaned
                    for token in _TERM_PATTERN.findall(sentence)
                    for cleaned in [_clean_term(token)]
                    if cleaned
                }
                for token in local_terms:
                    counts[token] += 1
                    sources_by_term.setdefault(token, set()).add(record.source_name)

        for term, frequency in counts.items():
            if frequency < 3:
                continue
            summary = self._find_best_summary(term, all_sentences)
            if not summary:
                continue
            draft = self._draft(term)
            draft.sources.update(sorted(sources_by_term.get(term, set()))[:4])
            draft.add_summary(summary, priority=90 + min(frequency, 20))
            for sentence in self._supporting_sentences(term, all_sentences):
                draft.add_scope(sentence)

    def _promote_canonical_bare_terms(self) -> None:
        snapshot = list(self.drafts.items())
        for name, draft in snapshot:
            evidence_pool = [name, draft.best_summary(), *draft.scope_parts[:4]]
            for evidence in evidence_pool:
                if not evidence:
                    continue
                for term, kind in _extract_canonical_bare_terms(evidence):
                    if not term:
                        continue
                    target = draft if term == name else self._draft(term)
                    target.sources.update(draft.sources)
                    target.categories.update(draft.categories)
                    if term != name:
                        target.aliases.add(name)
                        target.related.add(name)
                    target.related.update(list(draft.related)[:4])
                    current_summary = target.best_summary()
                    if not current_summary or _summary_quality_bonus(current_summary, term=term) < 2:
                        target.add_summary(
                            _canonical_bare_term_summary(term, kind),
                            priority=124,
                        )
                    if evidence != name:
                        target.add_scope(evidence)
                    for line in draft.scope_parts[:3]:
                        if term in line:
                            target.add_scope(line)

    def _merge_supporting_title_variants(self) -> None:
        changed = True
        while changed:
            changed = False
            for name in sorted(self.drafts, key=len, reverse=True):
                if name not in self.drafts:
                    continue
                parent_name = _canonical_parent_name(name, self.drafts.keys())
                if not parent_name or parent_name == name or parent_name not in self.drafts:
                    continue
                if _looks_like_queryable_artifact_alias(name):
                    continue
                self._merge_draft_into(parent_name, name)
                changed = True
                break

    def _merge_draft_into(self, parent_name: str, child_name: str) -> None:
        if parent_name == child_name or child_name not in self.drafts:
            return
        parent = self._draft(parent_name)
        child = self.drafts[child_name]
        child_summary = child.best_summary()

        parent.sources.update(child.sources)
        parent.categories.update(child.categories)
        parent.aliases.add(child.name)
        parent.aliases.update(child.aliases)
        parent.related.update(
            target for target in child.related if target not in {parent_name, child_name}
        )
        if child_summary:
            parent.add_scope(child_summary)
            if _summary_quality_bonus(child_summary, term=parent_name) > _summary_quality_bonus(
                parent.best_summary(),
                term=parent_name,
            ):
                parent.add_summary(child_summary, priority=174)
        for line in child.scope_parts:
            parent.add_scope(line)

        del self.drafts[child_name]
        for draft in self.drafts.values():
            if child_name in draft.related:
                draft.related.discard(child_name)
                if draft.name != parent_name:
                    draft.related.add(parent_name)

    def _fanout_structural_scope_lines(self) -> None:
        names = sorted(self.drafts, key=len, reverse=True)
        for name in names:
            draft = self.drafts.get(name)
            if draft is None:
                continue
            for line in list(draft.scope_parts):
                if not _is_structural_scope_line(line):
                    continue
                target_names: list[str] = []
                for other in names:
                    if other == name or other not in line:
                        continue
                    if other not in target_names:
                        target_names.append(other)
                for target_name in target_names:
                    if target_name == name or target_name not in self.drafts or target_name not in line:
                        continue
                    focused_line = _focused_clause_for_term(target_name, line)
                    if not focused_line or target_name not in focused_line:
                        continue
                    self.drafts[target_name].add_scope(focused_line)

    def _stabilize_summaries(self) -> None:
        for name, draft in self.drafts.items():
            current_summary = draft.best_summary()
            current_score = _summary_quality_bonus(current_summary, term=name) if current_summary else -99
            current_direct = _summary_has_direct_definition(name, current_summary)
            best_candidate = ""
            best_score = -99
            for line in draft.scope_parts:
                candidate = _summary_candidate_from_scope(name, line)
                if not candidate:
                    continue
                score = _summary_quality_bonus(candidate, term=name)
                if score > best_score or (score == best_score and len(candidate) > len(best_candidate)):
                    best_candidate = candidate
                    best_score = score
            if not best_candidate:
                continue
            candidate_direct = _summary_has_direct_definition(name, best_candidate)
            if current_summary and current_direct and not candidate_direct and current_score >= 0:
                continue
            if not current_summary or best_score >= current_score + 2 or len(current_summary) <= 18:
                draft.add_summary(best_candidate, priority=183 if not current_summary else 177)

    def _write_entries(self) -> None:
        self.entries_dir.mkdir(parents=True, exist_ok=True)
        self.placeholders_dir.mkdir(parents=True, exist_ok=True)
        for path in self.entries_dir.glob("*.md"):
            if path.name == ".gitkeep":
                continue
            path.unlink()

        valid_names = [
            name
            for name, draft in self.drafts.items()
            if draft.best_summary() and _draft_is_viable(name, draft)
        ]
        for name in sorted(valid_names):
            draft = self.drafts[name]
            aliases = sorted(
                alias
                for alias in draft.aliases
                if _should_keep_alias(alias, name=name, sources=draft.sources)
            )
            summary = _ensure_subject_sentence(name, draft.best_summary())
            scope_lines = _prioritize_scope_lines(name, draft.scope_parts, draft.related)
            related = _prioritize_related_targets(
                name,
                draft.related,
                scope_lines,
                valid_names=valid_names,
            )
            frontmatter = {
                "type": "concept",
                "status": "fact",
                "aliases": aliases,
                "sources": sorted(draft.sources)[:12],
            }
            lines = [
                "---",
                yaml.safe_dump(frontmatter, allow_unicode=True, sort_keys=False).strip(),
                "---",
                f"# {name}",
                "",
                summary,
                "",
                "## Scope",
                " ".join(scope_lines) or f"{name}用于描述相关系统中的关键概念与操作边界。",
            ]
            if related:
                lines.extend(["", "## Related"])
                for target in related:
                    lines.append(f"- [[{target}]] - 在同一运行场景中经常一起出现。")
            path = self.entries_dir / f"{name}.md"
            path.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")

    def _write_indexes(self) -> None:
        self.indexes_dir.mkdir(parents=True, exist_ok=True)
        for path in self.indexes_dir.glob("*.md"):
            path.unlink()

        categories: dict[str, list[str]] = {
            "core": [],
            "roles": [],
            "monitoring": [],
            "network": [],
            "safety": [],
            "operations": [],
            "transfer": [],
        }
        entry_names = sorted(
            path.stem for path in self.entries_dir.glob("*.md") if path.name != ".gitkeep"
        )
        for name in entry_names:
            draft = self.drafts[name]
            bucket = _primary_category(draft.categories, name)
            categories[bucket].append(name)

        root_links = []
        for category, names in categories.items():
            if not names:
                continue
            index_name = f"index.{category}"
            root_links.append(index_name)
            body = [
                "---",
                yaml.safe_dump(
                    {
                        "kind": "index",
                        "segment": category,
                        "last_tidied_at": date.today().isoformat(),
                        "entry_count": len(names),
                        "estimated_tokens": max(128, len(names) * 96),
                    },
                    allow_unicode=True,
                    sort_keys=False,
                ).strip(),
                "---",
                f"# {category} 索引",
                "",
                f"用于导航 {category} 相关的正式知识条目。",
                "",
            ]
            for name in names:
                body.append(f"- [[{name}]]")
            (self.indexes_dir / f"{index_name}.md").write_text(
                "\n".join(body).rstrip() + "\n",
                encoding="utf-8",
            )

        root_body = [
            "---",
            yaml.safe_dump(
                {
                    "kind": "index",
                    "segment": "root",
                    "last_tidied_at": date.today().isoformat(),
                    "entry_count": len(entry_names),
                    "estimated_tokens": max(256, len(entry_names) * 72),
                },
                allow_unicode=True,
                sort_keys=False,
            ).strip(),
            "---",
            "# 根索引",
            "",
            "白盒知识库的统一入口，优先把查询路由到概念、角色、监控、安全和网络等主题索引。",
            "",
        ]
        for target in root_links:
            root_body.append(f"- [[{target}]]")
        (self.kb_root / "index.root.md").write_text(
            "\n".join(root_body).rstrip() + "\n",
            encoding="utf-8",
        )

    def _draft(self, name: str) -> ConceptDraft:
        cleaned = _clean_term(name)
        if not cleaned:
            raise ValueError("concept name must not be empty")
        if cleaned not in self.drafts:
            self.drafts[cleaned] = ConceptDraft(name=cleaned)
        return self.drafts[cleaned]


def build_offline_batch(kb_root: str | Path, materials: Iterable[str | Path]) -> dict[str, Any]:
    builder = OfflineKnowledgeBuilder(kb_root)
    builder.load_existing_kb()
    return builder.ingest_materials(materials)


def tidy_offline_kb(kb_root: str | Path, *, focus: str = "general") -> dict[str, Any]:
    builder = OfflineKnowledgeBuilder(kb_root)
    builder.load_existing_kb()
    return builder.tidy(focus=focus)


def _load_source_record(path: Path) -> SourceRecord:
    text = _read_material_text(path)
    sentences = [
        _normalize_sentence(chunk)
        for chunk in _SENTENCE_PATTERN.split(text)
    ]
    sentences = [sentence for sentence in sentences if sentence]
    return SourceRecord(
        path=path,
        source_name=path.stem,
        text=text,
        sentences=sentences,
    )


def _read_material_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx" and DocxDocument is not None:
        doc = DocxDocument(path)
        return "\n".join(paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip())
    return path.read_text(encoding="utf-8", errors="ignore")


def _normalize_sentence(text: str) -> str:
    stripped = str(text).strip()
    if not stripped:
        return ""
    stripped = re.sub(r"^#+\s*", "", stripped)
    stripped = _strip_structural_prefix(stripped)
    stripped = re.sub(r"^\s*[-*>\d.\]\[()\s]+", "", stripped)
    stripped = re.sub(r"\[[ xX]\]\s*", "", stripped)
    stripped = stripped.replace("**", "").replace("`", "")
    stripped = re.sub(r"<[^>]+>", " ", stripped)
    stripped = re.sub(r"\bnote\s+(?:left|right|over)(?:\s+of)?[^:]*:", "", stripped, flags=re.IGNORECASE)
    stripped = re.sub(r"\s+", " ", stripped)
    stripped = stripped.replace("|", " ")
    stripped = re.sub(r"(\d)\s*[。．]\s*(\d)", r"\1.\2", stripped)
    stripped = re.sub(r"\s{2,}", " ", stripped).strip()
    if not stripped:
        return ""
    if len(stripped) > 280:
        stripped = stripped[:279].rstrip() + "。"
    if stripped[-1] not in "。！？!?":
        stripped += "。"
    return stripped


def _strip_contrastive_tail(text: str) -> str:
    cleaned = text
    for pattern in (
        r"[，,、]\s*(?:而)?不是[^。；;]*",
        r"[，,、]\s*并非[^。；;]*",
        r"[，,、]\s*不等于[^。；;]*",
    ):
        cleaned = re.sub(pattern, "", cleaned)
    return cleaned.strip()


def _clean_term(value: Any) -> str:
    text = str(value or "").strip()
    text = text.strip("`'\"<>[]（）()：:，,。！？!?；;")
    text = re.sub(r"\s+", "", text)
    if not _looks_like_concept_name(text) or _looks_like_fragmented_title(text):
        return ""
    return text


def _should_keep_alias(alias: str, *, name: str, sources: set[str]) -> bool:
    if not alias or alias == name or len(alias) > 48:
        return False
    if alias in sources and not _looks_like_queryable_artifact_alias(alias):
        return False
    if _looks_like_fragmented_title(alias):
        return False
    if not _alias_surface_compatible(alias, name=name):
        return False
    return True


def _looks_like_queryable_artifact_alias(text: str) -> bool:
    lowered = text.lower()
    return any(marker in lowered for marker in _QUERYABLE_ARTIFACT_ALIAS_MARKERS)


def _alias_surface_compatible(alias: str, *, name: str) -> bool:
    normalized_alias = _normalize_summary_surface(alias)
    normalized_name = _normalize_summary_surface(name)
    if not normalized_alias or not normalized_name:
        return False
    if normalized_alias == normalized_name:
        return True
    if normalized_alias in normalized_name or normalized_name in normalized_alias:
        return True
    if _looks_like_queryable_artifact_alias(alias):
        return True
    if re.fullmatch(r"[A-Za-z0-9_.-]+", alias):
        return True

    shared_prefix = 0
    for left, right in zip(normalized_alias, normalized_name):
        if left != right:
            break
        shared_prefix += 1
    if shared_prefix >= 3 and abs(len(normalized_alias) - len(normalized_name)) <= 8:
        return True

    shared_suffix = 0
    for left, right in zip(reversed(normalized_alias), reversed(normalized_name)):
        if left != right:
            break
        shared_suffix += 1
    if shared_suffix >= 3 and abs(len(normalized_alias) - len(normalized_name)) <= 8:
        return True

    return False


def _looks_like_concept_name(text: str) -> bool:
    if not text or len(text) < 2 or len(text) > 16:
        return False
    if text in _GENERIC_TERMS:
        return False
    if not re.search(r"[\u4e00-\u9fff]", text):
        return False
    if any(ch in text for ch in "#/\\"):
        return False
    if text.startswith("FILE") or text.startswith("http"):
        return False
    if text.startswith(("它", "该", "此", "这些", "那些")):
        return False
    if text.startswith(_BAD_TERM_PREFIXES):
        return False
    if text.endswith(_BAD_TERM_SUFFIXES):
        return False
    if any(ch.isdigit() for ch in text):
        return False
    if len(text) > 6 and any(marker in text for marker in ("必须", "需要", "需", "可以", "负责", "用于")):
        return False
    return bool(_TERM_PATTERN.fullmatch(text))


def _looks_like_fragmented_title(text: str) -> bool:
    compact = str(text or "").replace(" ", "")
    if not compact:
        return True
    if compact.endswith("职责"):
        return True
    if "中的一个" in compact:
        return True
    if any(
        marker in compact
        for marker in (
            "通过",
            "包括",
            "作为",
            "用于描述",
            "需在",
            "需与",
            "根据",
            "配备",
            "提交",
            "保持",
            "发现的",
            "启动晶格化",
            "系统运行中的一种",
        )
    ):
        return True
    if len(compact) > 4 and "是" in compact:
        return True
    if len(compact) > 5 and any(
        marker in compact
        for marker in (
            "内容",
            "现状",
            "情况",
            "方式",
            "模块",
            "参数",
            "日志",
            "记录",
            "说明",
            "职责",
            "时间同步",
            "评分细则",
        )
    ):
        return True
    if compact.startswith(("通过", "使用", "执行", "检查", "分析", "查看", "记录", "启动", "根据", "配备", "确认")):
        return True
    return False


def _ensure_subject_sentence(term: str, sentence: str) -> str:
    normalized = _focused_clause_for_term(term, sentence)
    if not normalized:
        return ""
    if normalized.startswith(term):
        quality_prefix = f"{term}数据质量"
        if normalized.startswith(quality_prefix):
            rest = normalized[len(quality_prefix):].lstrip("，,:： ")
            if rest:
                return f"{term}的数据质量{rest.rstrip('。')}。"
        return normalized
    for variant in _summary_term_variants(term):
        if variant != term and normalized.startswith(variant):
            return f"{term}{normalized[len(variant):]}"
    if term in normalized:
        return normalized
    if normalized.startswith(("负责", "用于", "定义", "描述", "记录", "监控", "建立", "新员工称呼", "校准")):
        return f"{term}是{normalized.rstrip('。')}。"
    return f"{term}是{normalized.rstrip('。')}。"


def _focused_clause_for_term(term: str, sentence: str) -> str:
    normalized = _normalize_sentence(sentence)
    if not normalized:
        return ""
    sentence_chunks = [
        chunk.strip()
        for chunk in re.split(r"(?<=[。！？!?])", normalized)
        if chunk.strip()
    ]
    if sentence_chunks:
        for chunk in sentence_chunks:
            if term in chunk:
                return chunk.rstrip("。") + "。"
    clauses = [
        clause.strip()
        for clause in re.split(r"[，；;]", normalized.rstrip("。"))
        if clause.strip()
    ]
    if not clauses:
        return normalized
    variants = _summary_term_variants(term)
    for clause in clauses:
        if any(clause.startswith(variant) for variant in variants):
            return clause.rstrip("。") + "。"
    for clause in clauses:
        if term in clause:
            return clause.rstrip("。") + "。"
    return normalized


def _structural_scope_summary(term: str, sentence: str) -> str:
    normalized = _focused_clause_for_term(term, sentence) or _normalize_sentence(sentence)
    if not normalized:
        return ""
    if normalized.startswith(f"{term}数据质量"):
        if "检查" in normalized:
            factors = normalized.split("检查", 1)[1].rstrip("。")
            return f"{term}是需要结合{factors}共同判断的数据质量指标。"
        return _ensure_subject_sentence(term, normalized)
    if f"{term}历史记录" in normalized and "完整性" in normalized:
        return f"{term}的历史记录完整性是判断相关运行质量的重要依据。"
    if normalized.startswith(f"{term}部署策略"):
        return _ensure_subject_sentence(
            term,
            normalized.replace(f"{term}部署策略", f"{term}的部署策略", 1),
        )
    if f"{term}的部署策略" in normalized:
        return _ensure_subject_sentence(term, normalized)
    for action in ("触发", "执行"):
        match = re.search(
            rf"(.+?)(?:时)?(?:直接|立即|提前|自动)?{action}{re.escape(term)}",
            normalized,
        )
        if not match:
            continue
        condition = match.group(1).strip("，,、 ")
        condition = re.sub(r"^(?:当|若|如|如果)", "", condition).strip()
        condition = re.sub(r"(?:时|后)$", "", condition).strip()
        if condition:
            return f"{term}的触发条件是{condition}。"
    return ""


def _summary_candidate_from_scope(term: str, sentence: str) -> str:
    normalized = _normalize_sentence(sentence)
    if not normalized or _line_is_noise(normalized):
        return ""
    structural = _structural_scope_summary(term, normalized)
    if structural:
        return structural
    if term in normalized and _summary_has_direct_definition(term, normalized):
        return _ensure_subject_sentence(term, normalized)
    if term in normalized and any(marker in normalized for marker in _SUMMARY_BOOST_MARKERS):
        return _ensure_subject_sentence(term, normalized)
    if re.search(rf"呈现{re.escape(term)}光泽", normalized):
        return f"{term}是镀层完好时表面呈现的健康光泽状态。"
    if re.search(rf"{re.escape(term)}(?:状态|阶段)", normalized) and any(
        marker in normalized for marker in ("光泽", "老化", "底噪", "峰谷差", "毛刺")
    ):
        return f"{term}是用于描述当前运行状态或健康阶段的正式标记。"
    return ""


def _summary_has_direct_definition(term: str, text: str) -> bool:
    normalized = _normalize_sentence(text)
    if not normalized:
        return False
    variants = _summary_term_variants(term)
    if any(
        normalized.startswith(f"{variant}{marker}")
        for variant in variants
        for marker in _DIRECT_DEFINITION_MARKERS
    ):
        return True
    return term in normalized and _has_direct_definition_signal(normalized)


def _is_generic_exception_label(label: str, subject: str) -> bool:
    compact = str(label or "").replace(" ", "")
    if not compact:
        return True
    if compact in {"基础", "操作", "运行", "控制"}:
        return True
    if subject and compact in {
        subject,
        f"{subject}基础",
        f"{subject}操作",
        f"{subject}控制",
        f"{subject}异常",
    }:
        return True
    return compact.endswith("基础") or compact.endswith("操作")


def _join_cn(values: Iterable[str]) -> str:
    parts = [str(item).strip() for item in values if str(item).strip()]
    return "、".join(dict.fromkeys(parts))


def _string_set(*values: Any) -> set[str]:
    result: set[str] = set()
    for value in values:
        text = str(value or "").strip()
        if text:
            result.add(text)
    return result


def _resolve_role_names(role_map: dict[str, Any], allowed_roles: list[str]) -> list[str]:
    result = []
    for role_key in allowed_roles:
        raw = role_map.get(role_key)
        if not isinstance(raw, dict):
            continue
        display_name = _clean_term(raw.get("display_name"))
        if display_name and display_name not in result:
            result.append(display_name)
    return result


def _category_for_domain(domain_name: str) -> str:
    mapping = {
        "system": "operations",
        "network": "network",
        "monitoring": "monitoring",
        "transfer": "transfer",
        "resonator": "operations",
        "maintenance": "operations",
    }
    return mapping.get(domain_name, "operations")


def _has_structural_signal(sentence: str) -> bool:
    return any(
        marker in sentence
        for marker in (
            "必须",
            "应",
            "需",
            "用于",
            "负责",
            "阈值",
            "告警",
            "权限",
            "步骤",
            "流程",
            "因为",
            "导致",
            "根因",
            "误判",
            "漏检",
            "换羽",
            "隔离",
        )
    )


def _prioritize_scope_lines(name: str, scope_parts: list[str], related: set[str]) -> list[str]:
    unique = []
    seen = set()
    for item in scope_parts:
        if not item or item in seen or _line_is_noise(item):
            continue
        seen.add(item)
        unique.append(item)

    def score(line: str) -> tuple[int, int]:
        return _line_signal_score(line, name=name, related=related), -len(line)

    ordered = sorted(unique, key=score, reverse=True)
    return ordered[:5]


def _prioritize_related_targets(
    name: str,
    related: set[str],
    scope_lines: list[str],
    *,
    valid_names: list[str],
) -> list[str]:
    valid_set = set(valid_names)

    def score(target: str) -> tuple[int, int]:
        mentions = sum(1 for line in scope_lines if target in line)
        causal = sum(
            1
            for line in scope_lines
            if target in line and any(marker in line for marker in _CAUSAL_MARKERS)
        )
        return causal * 3 + mentions, len(target)

    ordered = sorted(
        (
            target
            for target in related
            if target in valid_set and target != name and not _looks_like_fragmented_title(target)
        ),
        key=score,
        reverse=True,
    )
    return ordered[:8]


def _strip_structural_prefix(text: str) -> str:
    match = re.match(r"^(?:[A-Za-z_][A-Za-z0-9_]*|[一二三四五六七八九十]+)\s*:\s*(.+)$", text)
    if not match:
        return text
    return match.group(1).strip(" \"'")


def _has_definition_signal(text: str) -> bool:
    lowered = text.lower()
    return any(marker in text for marker in _DEFINITION_MARKERS) or "description" in lowered


def _has_direct_definition_signal(text: str) -> bool:
    return any(marker in text for marker in _DIRECT_DEFINITION_MARKERS)


def _extract_definition_subjects(text: str) -> list[str]:
    subjects: list[str] = []
    for chunk in re.split(r"[，；;]", text):
        segment = chunk.strip()
        if not segment:
            continue
        for marker in ("是", "指", "用于", "负责", "表示", "衡量", "属于", "作为"):
            prefix, found, _rest = segment.partition(marker)
            if not found:
                continue
            cleaned_prefix = prefix.replace("**", "").strip()
            cleaned_prefix = re.sub(r"(?:既|也)$", "", cleaned_prefix).strip()
            candidate = _clean_term(cleaned_prefix)
            if candidate:
                subjects.append(candidate)
    return list(dict.fromkeys(term for term in subjects if _looks_like_concept_name(term)))


def _extract_intro_subjects(text: str) -> list[str]:
    subjects: list[str] = []
    for pattern in (
        r"^([\u4e00-\u9fff]{2,8})(?:团队|技术|工具|算法|机制|现象|系统)",
        r"^([\u4e00-\u9fff]{2,8})特征[：:]",
        r"^([\u4e00-\u9fff]{2,8})数据质量",
    ):
        match = re.match(pattern, text)
        if not match:
            continue
        candidate = _clean_term(match.group(1))
        if candidate:
            subjects.append(candidate)
    return list(dict.fromkeys(subjects))


def _extract_embedded_compound_subjects(text: str) -> list[str]:
    subjects: list[str] = []
    for pattern, keep_full in (
        (r"(?:由|通知|联系|交由|与)([\u4e00-\u9fff]{2,8})(团队)", True),
        (r"(?:内置了|采用|使用|配备|支持|通过|调用)([\u4e00-\u9fff]{2,8})(算法|技术|工具|机制)", False),
        (r"(?:由)([\u4e00-\u9fff]{2,8})(承担|负责|主持)", False),
    ):
        for match in re.finditer(pattern, text):
            bare = _clean_term(match.group(1))
            if bare:
                subjects.append(bare)
            if keep_full:
                full = _clean_term("".join(match.groups()))
                if full:
                    subjects.append(full)
    return list(dict.fromkeys(subjects))


def _extract_structural_subjects(text: str) -> list[str]:
    subjects: list[str] = []
    normalized = _normalize_sentence(text).rstrip("。")
    if not normalized:
        return subjects
    for suffix in sorted(_STRUCTURAL_SUBJECT_SUFFIXES, key=len, reverse=True):
        patterns = (
            rf"([\u4e00-\u9fff]{{2,8}}){re.escape(suffix)}",
            rf"([\u4e00-\u9fff]{{2,8}})的{re.escape(suffix)}",
        )
        for pattern in patterns:
            for match in re.finditer(pattern, normalized):
                candidate = _clean_term(match.group(1))
                if candidate:
                    subjects.append(candidate)
    for pattern in _STRUCTURAL_SUPPORT_SUBJECT_PATTERNS:
        for match in pattern.finditer(normalized):
            candidate = _clean_term(match.group(1))
            if candidate:
                subjects.append(candidate)
    return list(dict.fromkeys(subjects))


def _extract_canonical_bare_terms(text: str) -> list[tuple[str, str]]:
    results: list[tuple[str, str]] = []
    for pattern, kind in _CANONICAL_BARE_TERM_PATTERNS:
        for match in pattern.finditer(text):
            candidate = _clean_term(match.group(1))
            if not candidate:
                continue
            results.append((candidate, kind))
    return results


def _extract_heading_terms(heading: str) -> list[str]:
    cleaned = heading.strip()
    cleaned = re.sub(r"^[一二三四五六七八九十\d.（）()\s-]+", "", cleaned).strip()
    if not cleaned:
        return []
    results: list[str] = []
    pair_match = re.match(r"^([\u4e00-\u9fff]{2,8})与([\u4e00-\u9fff]{2,8}).*$", cleaned)
    if pair_match:
        for value in pair_match.groups():
            candidate = _clean_term(value)
            if candidate:
                results.append(candidate)
    for suffix in _SECTION_TERM_SUFFIXES:
        if cleaned.endswith(suffix):
            candidate = _clean_term(cleaned[: -len(suffix)])
            if candidate:
                results.append(candidate)
    direct = _clean_term(cleaned)
    if direct and not results:
        results.append(direct)
    return list(dict.fromkeys(term for term in results if _looks_like_concept_name(term)))


def _document_subject_terms(record: SourceRecord) -> list[str]:
    subjects: list[str] = []
    for candidate in _extract_title_subjects(record.path.stem):
        if candidate not in subjects:
            subjects.append(candidate)
    for raw_line in record.text.splitlines():
        heading_match = re.match(r"^\s*#\s+(.+?)\s*$", raw_line.strip())
        if not heading_match:
            continue
        for candidate in _extract_title_subjects(heading_match.group(1)):
            if candidate not in subjects:
                subjects.append(candidate)
        break
    return subjects[:4]


def _extract_title_subjects(text: str) -> list[str]:
    cleaned = str(text or "").strip()
    if not cleaned:
        return []
    compact = cleaned.replace(" ", "")
    raw_compact = compact
    compact = re.sub(r"20\d{2}年Q\d", "", compact)
    compact = re.sub(r"20\d{2}年", "", compact)
    compact = re.sub(r"[（(][^）)]*[）)]", "", compact)
    compact = re.sub(r"[Vv]?\d+(?:\.\d+)*", "", compact)
    compact = compact.strip("-_:： ")
    results: list[str] = []
    for raw in re.findall(r"[（(]([^）)]+)[）)]", raw_compact):
        candidate = _clean_term(raw)
        if candidate:
            results.append(candidate)
    skip_direct = False
    for suffix, replacement in _ARTIFACT_SUBJECT_REWRITES:
        if not compact.endswith(suffix):
            continue
        prefix = compact[: -len(suffix)].strip("：:-_ ")
        candidate = _clean_term(prefix + replacement)
        if candidate:
            results.append(candidate)
            skip_direct = True
    for suffix in _DOC_SUBJECT_SUFFIXES:
        if suffix not in compact:
            continue
        skip_direct = True
        prefix = compact.split(suffix, 1)[0].strip("：:-_ ")
        candidate = _clean_term(prefix)
        if candidate:
            results.append(candidate)
        combined = _derive_doc_subject(prefix, suffix)
        if combined:
            results.append(combined)
        for trim_suffix in _DOC_SUBJECT_TRIM_SUFFIXES:
            if prefix.endswith(trim_suffix):
                bare = _clean_term(prefix[: -len(trim_suffix)])
                if bare:
                    results.append(bare)
        for bare_term, _kind in _extract_canonical_bare_terms(prefix):
            results.append(bare_term)
    direct = "" if skip_direct else _clean_term(compact)
    if direct:
        results.append(direct)
    return list(dict.fromkeys(term for term in results if _looks_like_concept_name(term)))


def _should_project_section_to_doc_subject(heading: str, informative_lines: list[str]) -> bool:
    heading_text = _normalize_sentence(heading).rstrip("。")
    if any(marker in heading_text for marker in _DOC_SECTION_MARKERS):
        return True
    return any(_has_structural_signal(line) for line in informative_lines[:4])


def _can_project_section_summary(term: str, heading: str, sentences: list[str]) -> bool:
    if not sentences or term in _extract_heading_terms(heading):
        return False
    heading_text = _normalize_sentence(heading).rstrip("。")
    if not any(marker in heading_text for marker in ("定义", "职责", "阶段", "规则", "评分")):
        return False
    return _has_definition_signal(sentences[0]) or _has_structural_signal(sentences[0])


def _derive_doc_subject(prefix: str, suffix: str) -> str:
    if not prefix:
        return ""
    for trailer in ("管理", "标准", "指南", "手册", "办法", "报告", "总结"):
        if suffix.endswith(trailer) and len(suffix) > len(trailer) + 1:
            candidate = _clean_term(prefix + suffix[: -len(trailer)])
            if candidate:
                return candidate
    return ""


def _best_heading_body_summary(term: str, sentences: list[str]) -> str:
    if not sentences:
        return ""
    for sentence in sentences:
        if term in sentence and _has_definition_signal(sentence):
            return sentence
    for sentence in sentences:
        if term in sentence:
            return _ensure_subject_sentence(term, sentence)
    return _ensure_subject_sentence(term, sentences[0])


def _select_informative_lines(
    lines: list[str],
    *,
    name: str,
    related: set[str],
    limit: int,
) -> list[str]:
    unique = []
    seen = set()
    for line in lines:
        if not line or line in seen or _line_is_noise(line):
            continue
        seen.add(line)
        unique.append(line)

    def score(line: str) -> tuple[int, int]:
        return _line_signal_score(line, name=name, related=related), -len(line)

    return sorted(unique, key=score, reverse=True)[:limit]


def _line_is_noise(text: str) -> bool:
    stripped = _normalize_sentence(text).rstrip("。")
    if not stripped:
        return True
    if any(stripped.endswith(suffix) for suffix in _DOC_SUBJECT_SUFFIXES):
        return True
    if any(f"{suffix}是" in stripped for suffix in _DOC_SUBJECT_SUFFIXES):
        return True
    if any(stripped.endswith(suffix) for suffix, _replacement in _ARTIFACT_SUBJECT_REWRITES):
        return True
    if "____" in stripped or "______" in stripped:
        return True
    if re.search(r"\b\d{1,2}:\d{2}:\d{2}\b", stripped):
        return True
    if re.search(r"\b\d+\s+\d+(?:\.\d+)?%\b", stripped):
        return True
    if stripped.startswith(("FILE:", "' ", '" ')):
        return True
    if any(marker in stripped for marker in ("-->", "└", "├", "│", "□", "@brief", "<description>", "模板编号")):
        return True
    if re.fullmatch(r"[0-9:：/.\-+ ]+", stripped):
        return True
    if stripped in {"步骤", "流程", "说明", "注意", "结果", "内容", "字段", "备注", "成功", "失败", "通过", "负"}:
        return True
    if re.search(r"(?:日期|时间|提交日期|生效日期)\s*[:：_＿]{1,}", stripped):
        return True
    if stripped.endswith(("：", ":")) and len(stripped) <= 12:
        return True
    return False


def _is_structural_scope_line(text: str) -> bool:
    normalized = _normalize_sentence(text)
    if not normalized or _line_is_noise(normalized):
        return False
    return any(marker in normalized for marker in _STRUCTURAL_SCOPE_MARKERS)


def _extract_xml_params(text: str) -> list[dict[str, str]]:
    results: list[dict[str, str]] = []
    for match in re.finditer(
        r'<param name="([^"]+)" value="([^"]+)"(?: unit="([^"]+)")?>(.*?)</param>',
        text,
        flags=re.DOTALL,
    ):
        body = match.group(4)
        description_match = re.search(r"<description>(.*?)</description>", body, flags=re.DOTALL)
        results.append(
            {
                "name": match.group(1),
                "value": match.group(2),
                "unit": match.group(3) or "",
                "description": _normalize_sentence(description_match.group(1)) if description_match else "",
            }
        )
    for match in re.finditer(r'<param name="([^"]+)" value="([^"]+)"(?: unit="([^"]+)")?\s*/>', text):
        results.append(
            {
                "name": match.group(1),
                "value": match.group(2),
                "unit": match.group(3) or "",
                "description": "",
            }
        )
    return results


def _extract_xml_rule_facts(text: str) -> list[dict[str, str]]:
    facts: list[dict[str, str]] = []
    for match in re.finditer(
        r"<rule\b.*?<condition field=\"([^\"]+)\" operator=\"([^\"]+)\" threshold=\"([^\"]+)\"[^>]*>(.*?)</condition>.*?<action type=\"([^\"]+)\"[^>]*>",
        text,
        flags=re.DOTALL,
    ):
        description_match = re.search(r"<description>(.*?)</description>", match.group(4), flags=re.DOTALL)
        facts.append(
            {
                "field": match.group(1),
                "operator": match.group(2),
                "threshold": match.group(3),
                "description": description_match.group(1) if description_match else "",
                "action": match.group(5),
            }
        )
    return facts


def _structured_param_subject(field_name: str) -> str:
    name = str(field_name or "").strip()
    for suffix in (
        "触发阈值",
        "触发条件",
        "部署策略",
        "计划周期",
        "执行周期",
        "覆盖率",
        "达标率",
        "补偿方式",
        "同步窗口",
        "分流比例",
    ):
        if not name.endswith(suffix) or len(name) <= len(suffix) + 1:
            continue
        candidate = _clean_term(name[: -len(suffix)])
        if candidate:
            return candidate
    return ""


def _structured_param_fact(
    *,
    subject: str,
    field_name: str,
    value: str,
    unit: str,
    description: str,
) -> str:
    unit_suffix = unit or ""
    if field_name.endswith(("触发阈值", "触发条件")):
        if description:
            return _normalize_sentence(description)
        return f"{subject}的{field_name}为{value}{unit_suffix}。"
    if field_name.endswith("部署策略"):
        return f"{subject}的部署策略为{value}{unit_suffix}。"
    if field_name.endswith(("计划周期", "执行周期", "同步窗口", "分流比例", "覆盖率", "达标率", "补偿方式")):
        return f"{subject}的{field_name}为{value}{unit_suffix}。"
    return ""


def _line_signal_score(
    line: str,
    *,
    name: str = "",
    related: Iterable[str] = (),
) -> int:
    value = 0
    if name and name in line:
        value += 2
    if any(marker in line for marker in _CAUSAL_MARKERS):
        value += 4
    if any(marker in line for marker in _SUMMARY_BOOST_MARKERS):
        value += 4
    if any(marker in line for marker in _HIGH_SIGNAL_SCOPE_MARKERS):
        value += 3
    if any(target in line for target in related):
        value += 1
    if any(marker in line for marker in ("步骤", "流程", "阈值", "周期", "窗口", "条件", "阶段")):
        value += 2
    if any(marker in line for marker in ("创建", "确认", "执行", "启动", "通知", "监测", "检查")):
        value += 2
    if re.search(
        r"\d+(?:\.\d+)?\s*(?:hz|%|小时|天|分钟|秒|毫秒|个月|米|单位|哈基米|tb|ms)",
        line.lower(),
    ):
        value += 3
    if len(line) >= 20:
        value += 1
    if any(marker in line for marker in ("```", "{", "}", "http", "./", "<description>")):
        value -= 3
    if _line_is_noise(line):
        value -= 6
    return value


def _extract_inline_fact_candidates(line: str) -> list[tuple[str, str]]:
    cleaned = str(line or "").strip()
    if not cleaned:
        return []
    cleaned = re.sub(r"^[\s>*#/\"'`-]+", "", cleaned)
    cleaned = cleaned.replace("**", "")
    candidates: list[tuple[str, str]] = []
    for pattern in (
        r"^([\u4e00-\u9fff]{2,12})\s*[：:]\s*(.+)$",
        r"^([\u4e00-\u9fff]{2,12})\s*[-—–]\s*(.+)$",
    ):
        match = re.match(pattern, cleaned)
        if not match:
            continue
        term = _clean_term(match.group(1))
        definition = _normalize_sentence(match.group(2))
        if not term or not definition or _line_is_noise(definition):
            continue
        if _looks_like_fragmented_title(term):
            continue
        if len(definition) < 6 and not _has_definition_signal(definition):
            continue
        candidates.append((term, _ensure_subject_sentence(term, definition)))
    return candidates


def _extract_bullet_fact_candidates(line: str) -> list[tuple[str, str]]:
    candidates: list[tuple[str, str]] = []
    normalized = line.replace("**", "").strip()
    for pattern in (
        r"^([\u4e00-\u9fff]{2,8})特征[：:](.+)$",
        r"^([\u4e00-\u9fff]{2,8})[：:](.+)$",
    ):
        match = re.match(pattern, normalized)
        if not match:
            continue
        term = _clean_term(match.group(1))
        summary = _normalize_sentence(match.group(2))
        if term and summary:
            candidates.append((term, _ensure_subject_sentence(term, summary)))
    return candidates


def _split_code_doc_title(first_line: str) -> tuple[str, str]:
    normalized = first_line.replace("—", "-").replace("–", "-").strip()
    if " - " not in normalized:
        return "", normalized
    left, right = normalized.split(" - ", 1)
    term = _clean_term(left)
    return term, right.strip()


def _extract_code_subject(text: str, fallback: str = "") -> str:
    for line in str(text).splitlines():
        cleaned = _normalize_sentence(line).rstrip("。")
        if not cleaned:
            continue
        for candidate in _extract_heading_terms(cleaned):
            if candidate:
                return candidate
        for candidate in _extract_intro_subjects(cleaned):
            if candidate:
                return candidate
    return _clean_term(fallback)


def _exception_label_from_doc(doc: str, class_name: str) -> str:
    normalized = _normalize_sentence(doc).rstrip("。")
    if normalized:
        cleaned = normalized.replace("异常", "")
        candidate = _clean_term(cleaned)
        if candidate:
            return candidate
    stem = class_name.replace("Error", "")
    return _clean_term(stem)


def _summary_term_variants(term: str) -> list[str]:
    variants = [term]
    for suffix in ("团队", "仪式", "流程", "阶段"):
        if term.endswith(suffix) and len(term) > len(suffix) + 1:
            variants.append(term[: -len(suffix)])
    if term.endswith("生命周期") and len(term) > 4:
        prefix = term[: -len("生命周期")]
        variants.append(f"{prefix}的完整生命周期")
    return list(dict.fromkeys(item for item in variants if item))


def _normalize_merge_title_core(text: str) -> str:
    normalized = _normalize_summary_surface(text)
    for marker in _MERGE_TITLE_WRAPPER_MARKERS:
        normalized = normalized.replace(_normalize_summary_surface(marker), "")
    return normalized


def _canonical_parent_name(name: str, existing_names: Iterable[str]) -> str:
    if _looks_like_queryable_artifact_alias(name):
        return ""
    candidates = [candidate for candidate in existing_names if candidate != name]
    if not candidates:
        return ""

    for suffix in sorted(_MERGEABLE_TITLE_SUFFIXES, key=len, reverse=True):
        if not name.endswith(suffix) or len(name) <= len(suffix) + 1:
            continue
        base = _clean_term(name[: -len(suffix)])
        if base and base in candidates:
            return base

    structured_suffix = next(
        (suffix for suffix in _MERGE_STRUCTURED_SUFFIXES if name.endswith(suffix)),
        "",
    )

    normalized_name = _normalize_summary_surface(name)
    if not normalized_name:
        return ""

    wrapper_markers = tuple(
        marker
        for marker in (_normalize_summary_surface(item) for item in _MERGE_TITLE_WRAPPER_MARKERS)
        if marker
    )
    child_core = _normalize_merge_title_core(name[: -len(structured_suffix)] if structured_suffix else name)
    best_name = ""
    best_score = (-1, -1)

    for candidate in candidates:
        normalized_candidate = _normalize_summary_surface(candidate)
        if (
            not normalized_candidate
            or normalized_candidate == normalized_name
            or len(normalized_candidate) >= len(normalized_name)
        ):
            continue

        score: tuple[int, int] | None = None
        if structured_suffix:
            if not candidate.endswith(structured_suffix):
                continue
            parent_core = _normalize_merge_title_core(candidate[: -len(structured_suffix)])
            if not child_core or not parent_core:
                continue
            if child_core == parent_core:
                score = (4, len(parent_core))
            elif parent_core in child_core:
                remainder = child_core.replace(parent_core, "", 1)
                if remainder and all(marker in remainder for marker in wrapper_markers):
                    score = (3, len(parent_core))
        else:
            candidate_core = _normalize_merge_title_core(candidate)
            if not candidate_core or candidate_core == child_core:
                if candidate_core == child_core and len(candidate) < len(name):
                    score = (3, len(candidate_core))
            elif candidate_core in child_core:
                remainder = child_core.replace(candidate_core, "", 1)
                if remainder and all(marker in remainder for marker in wrapper_markers):
                    score = (2, len(candidate_core))

        if score and score > best_score:
            best_name = candidate
            best_score = score

    return best_name


def _normalize_summary_surface(text: str) -> str:
    lowered = str(text or "").lower()
    for filler in ("的", "完整", "当前", "默认", "整体", "全系统", "全局"):
        lowered = lowered.replace(filler, "")
    return re.sub(r"[^a-z0-9\u4e00-\u9fff]+", "", lowered)


def _find_projected_record_summary(term: str, sentences: list[str]) -> str:
    for sentence in sentences:
        normalized = _normalize_sentence(sentence)
        if not normalized:
            continue
        if _line_is_noise(normalized) or re.match(r"^[一二三四五六七八九十\d]+[、.]", normalized):
            continue
        if _has_direct_definition_signal(normalized):
            return _ensure_subject_sentence(term, normalized)
        if any(marker in normalized for marker in ("阶段", "流程", "评分", "规则", "策略", "职责")):
            return _ensure_subject_sentence(term, normalized)
    return ""


def _find_compound_reference_summary(term: str, sentences: list[str]) -> str:
    variants = _summary_term_variants(term)
    for sentence in sentences:
        normalized = _normalize_sentence(sentence)
        if not normalized or _line_is_noise(normalized):
            continue
        for variant in variants:
            for suffix in ("算法", "技术", "工具", "机制", "团队", "系统", "协议"):
                if f"{variant}{suffix}" not in normalized:
                    continue
                if "负责" in normalized:
                    tail = normalized.split("负责", 1)[1].rstrip("。")
                    return f"{term}是负责{tail}的{suffix}。"
                if "可" in normalized:
                    tail = normalized.split("可", 1)[1].rstrip("。")
                    return f"{term}是用于{tail}的{suffix}。"
                return _ensure_subject_sentence(term, normalized)
    return ""


def _summary_quality_bonus(text: str, *, term: str = "") -> int:
    value = 0
    if any(
        marker in text
        for marker in (
            "阈值",
            "连续",
            "首次",
            "自动",
            "监测",
            "隔离",
            "传感器",
            "团队",
            "设备",
            "流程",
            "规则",
            "策略",
            "专家",
            "工程师",
            "角色",
            "衡量",
            "容器",
            "外壳",
            "死角",
            "区间",
            "比例",
        )
    ):
        value += 3
    if any(marker in text for marker in _SUMMARY_BOOST_MARKERS):
        value += 3
    if any(marker in text for marker in _CAUSAL_MARKERS):
        value += 2
    if any(marker in text for marker in _GENERIC_SUMMARY_MARKERS):
        value -= 2
    if 12 <= len(text) <= 80:
        value += 1
    value -= _low_information_summary_penalty(term, text)
    return value


def _draft_is_viable(name: str, draft: ConceptDraft) -> bool:
    summary = draft.best_summary()
    if not summary:
        return False
    if name in _GENERIC_TERMS:
        return False
    if not draft.sources:
        return False
    if _summary_quality_bonus(summary, term=name) <= -1 and len(draft.scope_parts) < 2:
        return False
    return True


def _is_stub_exception_doc(label: str, doc: str) -> bool:
    normalized = _normalize_sentence(doc).rstrip("。").replace(" ", "")
    if not normalized:
        return True
    return normalized in {
        label,
        f"{label}异常",
        f"{label}错误",
        f"{label}故障",
    }


def _low_information_summary_penalty(term: str, text: str) -> int:
    normalized = _normalize_sentence(text).rstrip("。")
    if not normalized:
        return 24

    compact = normalized.replace(" ", "")
    if "____" in normalized:
        return 18
    if re.search(r"(?:日期|时间|提交日期|生效日期)\s*[:：_＿-]{1,}", normalized):
        return 18
    if compact in {"日期", "时间", f"{term}时间", f"{term}日期"}:
        return 18

    if term:
        if compact in {
            term,
            f"{term}异常",
            f"{term}错误",
            f"{term}故障",
            f"{term}模块",
            f"{term}管理",
            f"{term}流程",
            f"{term}内容",
            f"{term}现状",
        }:
            return 24
        prefix = f"{term}是"
        if compact.startswith(prefix):
            tail = compact[len(prefix):]
            if tail.startswith(("检查", "读取", "执行", "查看", "核对", "通知", "记录", "获取", "统计", "评估", "判断", "确认")) and not any(
                marker in tail
                for marker in (
                    "角色",
                    "团队",
                    "工程师",
                    "专家",
                    "设备",
                    "工具",
                    "指标",
                    "阈值",
                    "状态",
                    "模式",
                    "规则",
                    "策略",
                    "协议",
                    "容器",
                    "外壳",
                    "死角",
                    "区间",
                    "比例",
                    "系统",
                    "工艺",
                    "流程",
                )
            ):
                return 16
    if compact.endswith(("模块", "管理", "流程", "内容", "现状")) and len(compact) <= max(8, len(term) + 4):
        return 12
    return 0


def _loaded_summary_priority(term: str, text: str) -> int:
    bonus = _summary_quality_bonus(text, term=term)
    if _low_information_summary_penalty(term, text) >= 16:
        return 96
    base = 156 if _has_definition_signal(text) else 144
    return base + max(-6, min(16, bonus * 2))


def _canonical_bare_term_summary(term: str, kind: str) -> str:
    if kind == "trace":
        return f"{term}是用于追踪来源、路径和传输记录的追踪机制。"
    if kind == "calibration":
        return f"{term}是用于频率设定、校准和协同稳定的设备或工具。"
    if kind == "metric":
        return f"{term}是用于衡量运行质量、稳定性或健康度的观测指标。"
    if kind == "sop":
        return f"{term}是需要按标准操作流程执行的正式操作或工具。"
    if kind == "mode":
        return f"{term}是系统运行中的一种模式或状态，用于表达当前处理阶段和稳定策略。"
    if kind == "initiative":
        return f"{term}是用于在特定窗口推进扩容、上线或全量运行的专项行动。"
    return f"{term}是正式流程中的一个阶段，通常承担关键参数注入、校准或切换准备。"


def _describe_schedule(raw: str) -> str:
    schedule = str(raw or "").strip()
    if not schedule:
        return "按既定节奏"
    if schedule in {"手动触发", "按需触发"}:
        return schedule
    parts = schedule.split()
    if len(parts) != 6:
        return schedule
    _sec, minute, hour, day, month, _week = parts
    if minute == "0" and hour.isdigit() and day == "*" and month == "*" and _week == "*":
        return f"每日{int(hour):02d}:00"
    if minute == "0" and hour.startswith("*/") and day == "*" and month == "*" and _week == "*":
        return f"每{hour[2:]}小时"
    if minute == "0" and hour == "0" and day == "1" and month.startswith("*/"):
        return f"每{month[2:]}个月"
    return schedule


def _metric_structured_summary(name: str, metric: dict[str, Any]) -> str:
    thresholds = metric.get("thresholds") or {}
    unit = str(metric.get("unit", "")).strip()
    unit_aliases = [str(item).strip() for item in metric.get("unit_aliases", []) if str(item).strip()]
    if name == "嗡鸣度":
        description = str(metric.get("description", "")).strip()
        resonance_peak = thresholds.get("resonance_peak") or {}
        lower = resonance_peak.get("min")
        upper = resonance_peak.get("max")
        if description and lower is not None and upper is not None and unit:
            return f"嗡鸣度是{description}，其理想工作区间通常为 {lower}-{upper}{unit}。"
        if description:
            return _ensure_subject_sentence(name, description)
    if name == "共振峰":
        lower = (thresholds.get("lower_bound") or {}).get("value")
        upper = (thresholds.get("upper_bound") or {}).get("value")
        if lower is not None and upper is not None and unit:
            return f"共振峰是嗡鸣度的理想运行区间，通常保持在 {lower}-{upper}{unit}。"
    if name == "清浊比":
        display_units = _join_cn([unit, *unit_aliases])
        if display_units:
            return f"清浊比是纯净哈基米与散斑的比例指标，通常以{display_units}表示。"
    if name == "饱和度":
        critical = (thresholds.get("critical") or {}).get("value")
        if critical is not None:
            suffix = "%" if unit == "percent" else unit
            return f"饱和度是谐振腔容量占最大容量的比例指标，超过 {critical}{suffix} 会进入临界处理区。"
    if name == "毛刺":
        warning = (thresholds.get("warning") or {}).get("value")
        if warning is not None:
            return f"毛刺是嗡鸣度瞬间异常的次数指标，短时间内超过 {warning} 次通常意味着存在设备干扰或异常波动。"
    return ""


def _route_type_summary(route_type: str) -> str:
    if route_type == "接力":
        return "接力是通过多个中继节点串联转发哈基米的路由方式。"
    if route_type == "分流":
        return "分流是把单条传输链路拆分为多条并行支路的路由方式。"
    if route_type == "合流":
        return "合流是把多条支路同步汇入同一目标腔体的路由方式。"
    return f"{route_type}是信使路由表中的一种正式路径选择方式。"


def _primary_category(categories: set[str], name: str) -> str:
    for category in ("core", "roles", "monitoring", "network", "safety", "operations", "transfer"):
        if category in categories:
            return category
    if any(marker in name for marker in ("掌灯人", "祭司团", "调音师", "守望者", "外乡人", "老把式", "守夜人", "编织者")):
        return "roles"
    if any(marker in name for marker in ("谐振腔", "哈基米", "晶格化", "清浊比")):
        return "core"
    if any(marker in name for marker in ("嗡鸣度", "盲区", "毛刺", "峰谷差", "回音壁")):
        return "monitoring"
    if any(marker in name for marker in ("潮涌", "红线", "锁龙井", "三振法则", "暗流")):
        return "safety"
    if any(marker in name for marker in ("织网", "驿站", "分流", "合流", "接力", "移星斗", "走线")):
        return "network"
    if any(marker in name for marker in ("旋涡协议", "跃迁", "断流", "回流", "剥离")):
        return "transfer"
    return "operations"
